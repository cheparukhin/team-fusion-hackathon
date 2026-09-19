"""Independently audit Psap:Lgals3 directly from pinned spreadsheets and references.

Uses explicit per-nucleotide exon coordinates, not the dataset reconstruction
coordinate mapper or the demo export. No inference about biological authenticity.
"""
import gzip
import hashlib
import json
import re
from collections import defaultdict
from datetime import datetime,timezone
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[2]
PAIR='Psap:Lgals3'


def run(root=ROOT):
    names=['data/raw/41586_2026_10982_MOESM5_ESM.xlsx',
           'data/raw/41586_2026_10982_MOESM10_ESM.xlsx',
           'data/reference/gencode.vM28.annotation.gtf.gz',
           'data/reference/gencode.vM28.transcripts.fa.gz']
    pinned={s['path']:s['sha256'] for s in json.loads((root/'results/dataset_reconstruction/manifest.json').read_text())['sources']}
    sources=[]
    for name in names:
        digest=hashlib.sha256((root/name).read_bytes()).hexdigest()
        if digest!=pinned[name]:raise ValueError(f'Source integrity mismatch: {name}')
        sources.append({'path':name,'sha256':digest})
    wb=openpyxl.load_workbook(root/names[0],read_only=True,data_only=True)
    ws=wb.active;rows=ws.iter_rows(values_only=True);headers=next(rows)
    reads=[]
    for index,values in enumerate(rows,2):
        row=dict(zip(headers,values))
        if row['Chimera_ID']==PAIR:
            reads.append({'spreadsheet_row':index,**row})
    wb.close()
    wb=openpyxl.load_workbook(root/names[1],read_only=True,data_only=True)
    probe=None;probe_row=None
    for index,row in enumerate(wb.active.iter_rows(values_only=True),1):
        if row[0]==PAIR:probe=row[2].upper();probe_row=index
    wb.close()
    if probe is None or len(probe)!=120:raise ValueError('Expected one 120-nt probe')
    exons=defaultdict(list)
    with gzip.open(root/names[2],'rt') as f:
        for line in f:
            if line.startswith('#'):continue
            cols=line.rstrip().split('\t')
            if cols[2]!='exon':continue
            attr=dict(re.findall(r'(\w+) "([^"]+)"',cols[8]))
            if attr.get('gene_name') not in PAIR.split(':'):continue
            exons[attr['transcript_id']].append((int(re.search(r'exon_number (\d+);',cols[8]).group(1)),cols[0],int(cols[3]),int(cols[4]),cols[6]))
    sequences={}
    with gzip.open(root/names[3],'rt') as f:
        tid=None
        for line in f:
            if line.startswith('>'):
                header=line[1:].strip().split('|');tid=header[0] if header[5] in PAIR.split(':') else None
                if tid:sequences[tid]=[header[5],'']
            elif tid:sequences[tid][1]+=line.strip()
    matches=[]
    for side,gene,half in [('a','Psap',probe[:60]),('b','Lgals3',probe[60:])]:
        for tid,(name,sequence) in sequences.items():
            if name!=gene:continue
            positions=[];chromosomes=set();strands=set()
            for _,chrom,start,end,strand in sorted(exons[tid]):
                chromosomes.add(chrom);strands.add(strand)
                positions.extend(range(start,end+1) if strand=='+' else range(end,start-1,-1))
            if len(positions)!=len(sequence):raise ValueError('Transcript/exon length mismatch')
            if len(chromosomes)!=1 or len(strands)!=1:raise ValueError('Inconsistent transcript')
            for offset in range(len(sequence)-len(half)+1):
                if sequence[offset:offset+len(half)]==half:
                    matches.append({'side':side,'gene':gene,'transcript_id':tid,'transcript_offset_zero_based':offset,
                                    'chromosome':next(iter(chromosomes)),'strand':next(iter(strands)),
                                    'terminal_genomic_base':positions[offset+59 if side=='a' else offset],
                                    'method':'Exact 60-nt match, explicitly expanded 1-based genomic exon coordinates'})
    endpoints={side:sorted({m['terminal_genomic_base'] for m in matches if m['side']==side}) for side in ['a','b']}
    if any(len(v)!=1 for v in endpoints.values()):raise ValueError('Ambiguous or missing probe coordinate')
    for row in reads:
        for side,letter in [('a','A'),('b','B')]:
            expected={(m['chromosome'],m['strand']) for m in matches if m['side']==side}
            if (row[f'Chromosome_Gene_{letter}'],row[f'Strand_Gene_{letter}']) not in expected:
                raise ValueError('Published read and probe locus/orientation mismatch')
    offsets=[{'source_row':row['spreadsheet_row'],'read_id':row['Read_ID'],
              'offset1':int(row['Breakpoint_Coordinate_Gene_A'])-endpoints['a'][0],
              'offset2':int(row['Breakpoint_Coordinate_Gene_B'])-endpoints['b'][0]} for row in reads]
    nearest=min(max(abs(o['offset1']),abs(o['offset2'])) for o in offsets)
    return {'status':'passed','verified_at':datetime.now(timezone.utc).isoformat(),'pair_id':PAIR,
            'sources':sources,'raw_read_rows':reads,'probe_spreadsheet_row':probe_row,'probe_sequence':probe,
            'independent_transcript_matches':matches,'probe_endpoints':endpoints,'read_offsets':offsets,
            'unique_read_ids':len({row['Read_ID'] for row in reads}),
            'closest_max_endpoint_difference_nt':nearest,
            'conclusion':f'The {nearest}-nt nearest discrepancy is reproduced directly from the pinned source spreadsheets and M28 reference, independently of the demo and reconstruction mapper.',
            'limits':['Published long-read base-origin convention remains unspecified; the discrepancy is descriptive.','This does not establish a mapping error, false RNA, probe failure, or alternative isoform. Original alignments/sample context require review.']}

if __name__=='__main__':
    result=run()
    path=ROOT/'results/review/psap_source_audit.json'
    path.write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps({k:result[k] for k in ['status','unique_read_ids','probe_endpoints','closest_max_endpoint_difference_nt','conclusion']},indent=2))
