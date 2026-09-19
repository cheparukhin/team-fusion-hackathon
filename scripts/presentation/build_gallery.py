"""Editable chRNA gallery deck plus shared-source SVG/PNG/PDF figure exports.
Authoring dependencies are optional and isolated from the analytical environment.
"""
from __future__ import annotations
import base64,csv,hashlib,html,json,math,textwrap
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR,MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
import cairosvg
from PIL import Image,ImageOps,ImageDraw,ImageFont
from pypdf import PdfReader,PdfWriter
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'results/presentation';G=OUT/'gallery';G.mkdir(parents=True,exist_ok=True)
B='#0000FF';R='#D22D27';INK='#172431';GRAY='#6B747E';LIGHT='#EEF1F6';MID='#CAD1DC';WHITE='#FFFFFF';PALE_B='#EBEDFF';PALE_R='#FBECEA'
W,H=1600,900
MET=json.loads((ROOT/'results/classifier/metrics.json').read_text());EX=json.loads((OUT/'candidate_examples.json').read_text());PILOT=json.loads((ROOT/'results/compute/pilot_summary.json').read_text());DATA=json.loads((ROOT/'demo/data.json').read_text())
C={c['pair_id']:c for c in DATA['candidates']};E={c['pair_id']:c for c in EX['candidates']}
# This is a reviewed frozen-result deck, not a generic template: fail if its narrative needs updating.
assert MET['cohort']['n']==479 and MET['cohort']['positives']==109 and MET['cohort']['groups']==352
assert MET['matched_hic']['rna']['n']==401 and MET['matched_hic']['rna']['positives']==92
assert PILOT['read_pairs']==2000000 and PILOT['n_chimeric_junction_records']==28482
assert PILOT['n_split_junction_records']==3275 and PILOT['matched_pairs']==0
assert abs(PILOT['wall_time_seconds']-85.49)<1e-9

PAPER='Venezia et al., Nature (2026), doi:10.1038/s41586-026-10982-x'
SLIDES=[]
class Scene:
 def __init__(self,slug,title,subtitle,source,method,caveat):
  self.slug,self.title,self.source,self.method,self.caveat=slug,title,source,method,caveat;self.items=[]
  self.rect(0,0,W,H,WHITE);self.rect(0,0,1000,7,B);self.rect(1000,0,600,7,R)
  self.text(64,35,1300,30,'chRNA / EVIDENCE GALLERY',18,GRAY,True)
  self.text(64,86,1472,78,title,45,INK,True)
  self.text(64,166,1472,62,subtitle,23,GRAY)
  self.line(64,813,1536,813,MID,1)
  self.text(64,829,1380,31,'SOURCE  '+source,13,GRAY)
  self.text(64,861,1390,23,'METHOD / LIMIT  '+caveat,12,GRAY)
  self.text(1465,838,70,30,f'{len(SLIDES)+1:02}',20,GRAY,align='right');SLIDES.append(self)
 def rect(self,x,y,w,h,fill=LIGHT,stroke=None,r=0,sw=1):self.items.append(dict(kind='rect',x=x,y=y,w=w,h=h,fill=fill,stroke=stroke,r=r,sw=sw))
 def line(self,x,y,x2,y2,color=INK,sw=2,dash=False):self.items.append(dict(kind='line',x=x,y=y,x2=x2,y2=y2,color=color,sw=sw,dash=dash))
 def circle(self,x,y,d,fill,stroke=None,sw=1):self.items.append(dict(kind='circle',x=x,y=y,w=d,h=d,fill=fill,stroke=stroke,sw=sw))
 def text(self,x,y,w,h,t,size=25,color=INK,bold=False,align='left',italic=False):
  # Measure actual glyph widths; share the same explicit breaks with PowerPoint and SVG.
  font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans'+('-Bold' if bold else '')+'.ttf'
  while True:
   font=ImageFont.truetype(font_path,max(1,int(size)));lines=[]
   for para in str(t).split('\n'):
    line=''
    for word in para.split():
     trial=(line+' '+word).strip()
     if line and font.getlength(trial)>w:lines.append(line);line=word
     else:line=trial
    lines.append(line)
   needed=size+(len(lines)-1)*size*1.18
   if needed<=h and max((font.getlength(l) for l in lines),default=0)<=w:break
   if size<=11:break
   size-=1
  self.items.append(dict(kind='text',x=x,y=y,w=w,h=h,text='\n'.join(lines),size=size,color=color,bold=bold,align=align,italic=italic))
 def arrow(self,x,y,x2,y2,color=INK,sw=3):
  self.line(x,y,x2,y2,color,sw);a=math.atan2(y2-y,x2-x);L=12
  for da in [-.48,.48]:self.line(x2,y2,x2-L*math.cos(a+da),y2-L*math.sin(a+da),color,sw)
 def image(self,path,x,y,w,h):
  path=Path(path)
  if not path.exists():raise FileNotFoundError(path)
  iw,ih=Image.open(path).size;factor=min(w/iw,h/ih);dw,dh=iw*factor,ih*factor
  self.items.append(dict(kind='image',path=str(path),x=x+(w-dw)/2,y=y+(h-dh)/2,w=dw,h=dh))
 def label(self,x,y,w,t,color=B):self.rect(x,y,w,38,PALE_B if color==B else PALE_R);self.text(x+12,y+7,w-24,25,t,17,color,True)
 def card(self,x,y,w,h,title,value,detail,color=B):
  self.rect(x,y,w,h,WHITE,MID,12);self.rect(x,y,5,h,color)
  self.text(x+22,y+20,w-42,40,title,20,GRAY,True);self.text(x+22,y+65,w-42,78,value,52,color,True);self.text(x+22,y+151,w-42,h-160,detail,20,GRAY)
 def notes(self):return f'{self.title}\n\nSOURCE\n{self.source}\n\nMETHOD\n{self.method}\n\nINTERPRETATION / CAVEAT\n{self.caveat}\n\nGenerated from the reviewed repository artifacts, not invented measurements. Exact sources and checksums are in gallery/figure_manifest.json.'

def exon_row(s,x,y,w,exons,color,label,muted=False):
 n=len(exons);gap=12;ew=min(105,(w-(n-1)*gap)/n)
 s.text(x,y-44,w,32,label,23,color,True,italic=True);s.line(x,y+26,x+w,y+26,MID,3)
 for i,nr in enumerate(exons):
  xx=x+i*(ew+gap);s.rect(xx,y,ew,54,LIGHT if muted else color,WHITE,0)
  s.text(xx,y+12,ew,32,str(nr),22,GRAY if muted else WHITE,True,align='center')
 s.arrow(x+w-8,y+26,x+w+20,y+26,color,2)

def genomic_note(s,c,x,y,w):
 p=c['probe_examples'][0];a=p['sides']['a']['representative'];b=p['sides']['b']['representative']
 s.text(x,y,w,54,f"{a['chromosome']}:{a['breakpoint']:,} ({a['strand']})  →  {b['chromosome']}:{b['breakpoint']:,} ({b['strand']})",21,GRAY)
 s.text(x,y+53,w,65,f"{a['transcript_id']}  /  {b['transcript_id']}",17,GRAY)

# 01 — Workflow first, as requested.
s=Scene('01_workflow','From published junctions to testable evidence','A reproducible workflow for mouse macrophage chimeric RNA candidates.',
 'Dataset audit + classifier MODEL_CARD.md + compute pilot_summary.json',
 'Published tables define the candidate panel and reporting target; gene-disjoint models use only prespecified RNA/spatial features. GPU and language-model outputs remain separate displayed evidence.',
 'Conceptual workflow; arrows describe analysis dependencies, not new experimental validation.')
steps=[('01','RECONSTRUCT','Tables 3 / 4 / 7 / 8\nGENCODE M28',B),('02','PRESERVE','Ordered parents\nExact junctions + reads',B),('03','MODEL','RNA features\n500-kb Hi-C context',R),('04','EVALUATE','5 grouped folds\nMatched cohorts',B),('05','EXPLORE','Rank + inspect\nTrace every claim',R)]
for i,(n,title,body,col) in enumerate(steps):
 x=65+i*305;s.rect(x,294,260,240,WHITE,MID,12);s.circle(x+20,314,50,col);s.text(x+20,324,50,30,n,24,WHITE,True,align='center');s.text(x+20,384,224,35,title,25,col,True);s.text(x+20,434,223,85,body,24,INK)
 if i<4:s.arrow(x+270,414,x+293,414,GRAY,3)
s.rect(65,603,720,133,PALE_B);s.text(88,623,680,35,'NVIDIA · independent evidence',27,B,True);s.text(88,669,680,52,'Bounded Parabricks pilot → exact junction matching → display only',23,INK)
s.rect(814,603,720,133,PALE_R);s.text(837,623,680,35,'OpenAI · source-grounded explanation',27,R,True);s.text(837,669,680,52,'Evidence rows + paper passages → reviewed cached reports',23,INK)

# 02 — cohort.
s=Scene('02_cohort','The evaluation unit is an ordered gene pair','529 probe designs do not mean 529 independently validated chimeras.',
 'dataset_reconstruction/audit.md; probe_panel.tsv; model_input.tsv',
 'Sequence-backed pair identity reconciliation; collapse repeated designs to ordered pairs; conservatively exclude unresolved candidate/control identities.',
 'NanoString membership is reported support. Unknown testing/QC is not a verified negative.')
for i,(v,t,d) in enumerate([('529','probe designs','Published design panel'),('527','ordered pairs','Two repeated-design collapses'),('479','eligible pairs','48 unresolved pairs excluded')]):
 x=65+i*505;s.card(x,266,455,235,t,v,d,B if i!=1 else R)
 if i<2:s.arrow(x+464,377,x+489,377,GRAY)
s.text(66,555,1200,45,'479 eligible pairs: one label per ordered pair',29,INK,True)
barx,bary,barw=66,628,1469;frac=109/479;s.rect(barx,bary,barw*frac,68,B);s.rect(barx+barw*frac,bary,barw*(1-frac),68,PALE_R)
s.text(80,708,600,55,'109 reported supported',30,B,True);s.text(735,708,800,55,'370 not reported supported',30,R,True,align='right')

# 03 evidence unit.
s=Scene('03_evidence_units','Different evidence answers different questions','Keep read identity, junction identity, pair-level reporting, and protein function separate.',
 PAPER+'; dataset_reconstruction/{junctions,probe_panel}.tsv',
 'Separate junction records, probe records, gene-pair target, and paper experiments. Unknown assay/QC and sample identifiers remain explicit.',
 'Do not promote a pair-level label to isoform validation, translation, function, or authenticity.')
for x,color,title,question in [(65,B,'READ / JUNCTION','What molecule or breakpoint was observed?'),(567,R,'ORDERED PAIR','Was the pair reported NanoString-supported?'),(1069,GRAY,'FUNCTION','Was a specific biological effect demonstrated?')]:
 s.rect(x,273,466,406,WHITE,MID,12);s.text(x+24,300,418,58,title,30,color,True);s.text(x+24,591,418,70,question,25,INK)
# schematic read stacks
for i in range(4):s.line(100,398+i*29,259+i*9,398+i*29,B,8);s.line(262+i*9,398+i*29,466,398+i*29,R,8)
s.line(270,367,270,541,GRAY,2,True);s.text(103,539,365,35,'Exact coordinates + read IDs',20,GRAY)
s.circle(633,405,96,B);s.circle(846,405,96,R);s.text(633,428,96,50,'A',42,WHITE,True,align='center');s.text(846,428,96,50,'B',42,WHITE,True,align='center');s.arrow(743,452,832,452,INK,3)
s.text(609,531,390,35,'Reported ≠ every isoform validated',19,GRAY,align='center')
s.text(1113,401,370,109,'?',84,GRAY,True,align='center');s.text(1102,531,395,42,'Requires separate experiments',20,GRAY,align='center')
s.label(65,718,1470,'Testing/QC unknown  •  Biological sample count unavailable  •  Missing contact ≠ zero contact',R)

# 04 exact Gsdmd architecture.
c=E['Gsdmd:Tmem106a'];a=c['probe_examples'][0]['sides']['a']['representative'];b=c['probe_examples'][0]['sides']['b']['representative']
s=Scene('04_gsdmd_architecture','Gsdmd → Tmem106a: RNA architecture and a predicted fold','Reference-reconstructed RNA and protein sequence; the candidate fold remains low confidence.',
 PAPER+' Figs 1 / 3; candidate_examples.json; reference_orf_verification.json; structure manifest',
 'Exact probe halves map to annotated transcripts. Exon numbers are transcript-specific; a lexicographic representative is displayed among exact matches. The 118-aa sequence was independently reference-reconstructed and matches the cached Boltz2 input.',
 'Symbolic exon widths. Retained RNA exons do not imply translation throughout. Boltz2 fold is predicted, not experimental.')
exon_row(s,75,321,365,[1,2],B,'Gsdmd · exons 1–2')
exon_row(s,535,321,415,[6,7,8,9],R,'Tmem106a · exons 6–9')
s.arrow(462,348,514,348,INK,4);s.text(455,291,70,42,'2 → 6',18,INK,True,align='center')
genomic_note(s,c,75,421,875)
s.text(76,562,875,43,'118-aa reconstructed protein architecture',29,INK,True)
start,width=76,872;cut=width*73/118;s.rect(start,632,cut,66,B);s.rect(start+cut,632,width-cut,66,R)
s.text(start+16,648,cut-32,37,'1–73 · Gsdmd-derived',25,WHITE,True)
s.text(start+cut+12,648,width-cut-24,37,'74–118 · novel C terminus',23,WHITE,True)
s.text(76,719,875,66,'Blue: in-frame Gsdmd sequence. Red: novel out-of-frame Tmem106a-derived C terminus.',24,GRAY)
s.line(991,272,991,773,MID,2)
s.text(1040,271,478,44,'Candidate fold · predicted',29,R,True)
s.image(ROOT/'results/structures/gsdmd_tmem106a/cartoon.png',1020,329,510,327)
s.text(1020,691,510,44,'Low confidence · mean pLDDT 48.70',24,R,True,align='center')
s.text(1020,741,510,39,'Boltz2 · pTM 0.336',23,GRAY,align='center')

# 05 two exact RNA-only examples.
s=Scene('05_candidate_contrasts','Junction evidence and model score can disagree','Exact exon choices are transcript-specific; a high score does not create biological support.',
 'candidate_examples.json; predictions.tsv; '+PAPER+' Fig. 1 / Extended Data Fig. 2',
 'Two probe-sequence-derived junction diagrams; same preserved parent order and assembly. Full-panel out-of-fold RNA scores shown; no ORF inference for these examples.',
 'Representative transcripts are not uniquely validated isoforms. No protein/function claim for Psap:Lgals3.')
for yi,pid in [(292,'Cd274:Lacc1'),(550,'Psap:Lgals3')]:
 c=E[pid];row=C[pid];p=c['probe_examples'][0];a=p['sides']['a']['representative'];b=p['sides']['b']['representative'];ga,gb=pid.split(':')
 s.text(66,yi-40,650,43,pid.replace(':',' → '),31,INK,True)
 exa=a['retained_exon_numbers'];exa=exa if len(exa)<=6 else ['…',*exa[-4:]]
 exon_row(s,66,yi+63,510,exa,B,ga);exon_row(s,755,yi+63,470,b['retained_exon_numbers'],R,gb);s.arrow(590,yi+89,713,yi+89,INK,3)
 s.text(602,yi+38,125,34,f"{a['junction_exon']} → {b['junction_exon']}",22,INK,True,align='center')
 s.text(1260,yi+13,270,54,f"{row['score_rna']:.3f}",43,B,True);s.text(1260,yi+67,270,36,'RNA OOF score',19,GRAY)
 s.text(66,yi+147,1140,45,f"{a['chromosome']}:{a['breakpoint']:,} ({a['strand']}) → {b['chromosome']}:{b['breakpoint']:,} ({b['strand']})",20,GRAY)
 s.text(1260,yi+116,270,60,'Reported support' if row['label']==1 else 'Not reported supported',21,B if row['label']==1 else R,True)

# 06 validation split.
s=Scene('06_leakage_control','Evaluate new parent-gene groups, not memorized genes','352 parent-gene components are held intact across 5 folds.',
 'classifier/folds.tsv; predictions.tsv; MODEL_CARD.md',
 'Connected components group every pair sharing either parent gene. Deterministic five-fold predictions (seed 42); fold-local median imputation and scaling; fixed L2 logistic regression C=1.',
 'Illustrative partition diagram; exact assignments are in folds.tsv. No validation-derived predictor inputs.')
for f in range(5):
 x=69+f*295;s.rect(x,288,258,192,PALE_B if f<4 else PALE_R,None,8);s.text(x+18,304,222,34,f'Fold {f}',25,B if f<4 else R,True)
 for q in range(3):
  xx=x+49+q*60;s.circle(xx,374,28,B if f<4 else R)
  if q<2:s.line(xx+28,388,xx+60,388,B if f<4 else R,2)
 s.text(x+16,433,226,28,'Training groups' if f<4 else 'Held-out groups',19,GRAY)
s.text(68,510,1030,40,'Fit on training rows only',30,B,True);s.text(1215,510,320,40,'Score once',30,R,True)
for i,txt in enumerate(['Median imputation\n+ missing indicators','Feature scaling','Logistic regression\nL2 · C = 1']):
 x=70+i*365;s.rect(x,571,325,102,WHITE,MID,8);s.text(x+19,590,290,66,txt,24,INK,True,align='center')
 if i<2:s.arrow(x+334,621,x+355,621,GRAY)
s.arrow(1127,621,1210,621,INK);s.rect(1224,571,310,102,PALE_R);s.text(1238,591,282,65,'Out-of-fold\nreported-support score',23,R,True,align='center')
s.label(68,718,1466,'Hi-C ablation: both classifiers use identical contact-observed training and test candidates.',B)

# 07 AP bars + difference interval.
s=Scene('07_average_precision','A small AP increase; uncertainty includes no benefit','Matched cohort: 401 ordered pairs, 92 reported supported. Identical training rows and fixed folds.',
 'classifier/metrics.json → matched_hic + comparison; predictions.tsv',
 'Average precision of saved out-of-fold predictions. Difference interval from 1,000 bootstrap resamples of parent-gene components, without model refitting.',
 'RNA AP 0.296 versus RNA+Hi-C AP 0.300 does not establish improvement; bootstrap interval crosses zero.')
m=MET['matched_hic'];chartx,charttop,chartbottom=115,292,650
for val in [0,.1,.2,.3,.4]:
 yy=chartbottom-val/.4*(chartbottom-charttop);s.line(100,yy,970,yy,LIGHT,2);s.text(55,yy-13,42,30,f'{val:.1f}',18,GRAY,align='right')
for i,(key,label,col) in enumerate([('read_support','Read support',GRAY),('rna','RNA',B),('hic','RNA + Hi-C',R)]):
 v=m[key]['average_precision'];x=160+i*280;h=v/.4*(chartbottom-charttop);s.rect(x,chartbottom-h,160,h,col);s.text(x-12,chartbottom-h-45,184,36,f'{v:.3f}',31,col,True,align='center');s.text(x-30,677,220,52,label,24,INK,True,align='center')
s.text(1035,290,480,40,'Hi-C − RNA',32,INK,True);s.text(1035,343,480,71,f"+{m['comparison']['delta_average_precision']:.4f}",56,R,True)
s.text(1035,423,480,65,'difference in average precision',24,GRAY)
lo,hi=m['comparison']['ci95'];xmin,xmax=1045,1495;axis=lambda v:xmin+(v+.05)/.11*(xmax-xmin);yy=593
s.line(xmin,yy,xmax,yy,MID,2);s.line(axis(0),yy-70,axis(0),yy+70,GRAY,2,True);s.line(axis(lo),yy,axis(hi),yy,R,6);s.circle(axis(m['comparison']['delta_average_precision'])-10,yy-10,20,R)
s.text(1019,689,498,58,f'95% interval [{lo:.3f}, {hi:.3f}]',25,INK,True,align='center')

# 08 top20 and subset.
s=Scene('08_top20','Top-ranked candidates tell a more cautious story','A small AP change does not guarantee better prioritization at a chosen cutoff.',
 'classifier/metrics.json → matched_hic / interchromosomal',
 'Counts of reported-supported pairs among the first 20 saved scores, with pair-ID tie-breaking. Interchromosomal comparison restricted to the same 244 observed-contact pairs.',
 'Dots represent reporting labels, not verified true/false chimeras. The compared ranked lists need not contain the same pairs.')
for row,(key,label,col) in enumerate([('read_support','Read-support ranking',GRAY),('rna','RNA',B),('hic','RNA + Hi-C',R)]):
 y=289+row*142;n=m[key]['supported_at_k'];s.text(67,y+7,325,65,label,28,col,True)
 for i in range(20):s.circle(425+i*42,y,27,col if i<n else LIGHT,MID if i>=n else None)
 s.text(1310,y-1,220,54,f'{n} / 20',35,col,True)
s.rect(65,710,1470,68,PALE_B);s.text(85,725,1425,42,'Interchromosomal matched subset: RNA 5 / 20  →  RNA + Hi-C 7 / 20   (244 pairs; 58 reported supported)',23,B,True)

# 09 actual NVIDIA.
s=Scene('09_nvidia_pilot','An actual NVIDIA run, with no probe-panel matches','Parabricks on an A100-SXM4-80GB, using a deterministic two-million-pair RNA subsample.',
 'compute/pilot_summary.json; evidence.tsv; Chimeric.out.junction; independent review',
 'Parabricks 4.7.1-1; paired 151-nt reads from SRR37513722; both ordered breakpoints, strands and chromosomes matched within ±10 nt to probe-derived junctions.',
 'Zero probe matches does not refute candidates. 85.49 s is alignment only; no matched CPU speedup benchmark.')
s.card(65,275,445,237,'VALIDATED READ PAIRS','2,000,000','Deterministic prefix subsample',B);s.card(577,275,445,237,'ALIGNMENT WALL TIME','85.49 s','Reference / setup excluded',R);s.card(1090,275,445,237,'MATCHED PROBE PAIRS','0','Independent GPU evidence only',GRAY)
s.text(69,563,700,50,'28,482 raw chimeric records',34,INK,True)
s.rect(69,641,850*(3275/28482),56,B);s.rect(69+850*(3275/28482),641,850*(25207/28482),56,LIGHT)
s.text(70,711,480,65,'3,275 resolved split-junction records',25,B,True);s.text(506,711,478,65,'25,207 encompassing-mate records',25,GRAY,True)
s.arrow(985,642,1085,642,GRAY);s.text(1131,579,393,95,'±10 nt\nordered-junction match',29,INK,True,align='center');s.text(1100,706,433,69,'$1.35 quote-based instance estimate\nInstance deleted; watchdog cancelled',20,GRAY,align='center')

# 10 reports.
s=Scene('10_openai_reports','OpenAI turns cached evidence into traceable explanations','Five Codex-authored reports separate observations, interpretations, and unknowns.',
 'demo/reports/*.json; sources/passages.json; report_review.md; scripts/demo/reports.py',
 'The Codex agent authored five reports from actual evidence rows and supplied primary-paper passages. Citation IDs, numerical claims, input snapshots and hashes were reviewed/cached.',
 'No Responses API call occurred without credentials. The optional strict-schema API adapter is separate; reports are not model features or labels.')
for x,t,col in [(66,'Evidence row',B),(434,'Primary passages',R)]:
 s.rect(x,293,310,178,WHITE,MID,10);s.text(x+23,313,264,46,t,29,col,True)
 for k in range(3):s.line(x+25,380+k*23,x+261-k*25,380+k*23,MID,7)
s.arrow(769,384,886,384,INK,4)
s.rect(916,275,616,419,WHITE,MID,10)
for j,(kind,claim,col) in enumerate([('OBSERVATION','Reported NanoString pair support',B),('INTERPRETATION','Paper CTCF perturbation context',R),('UNKNOWN','Assay QC; function beyond evidence',GRAY)]):
 y=307+j*125;s.label(943,y,205,kind,col if col!=GRAY else B);s.text(943,y+51,554,63,claim,25,INK)
s.text(66,539,754,66,'Source identity + exact input hash',33,INK,True);s.text(66,623,758,97,'Supported: Gsdmd:Tmem106a, Cd274:Lacc1\nContrast: Psap:Lgals3 is not reported supported',26,GRAY)
s.label(65,741,1467,'Observed evidence stays distinct from biological authenticity, translation, and function.',R)

# 11 structure assets are provided independently; never substitute fabricated coordinates.
sp=ROOT/'results/structures/manifest.json';sm=json.loads(sp.read_text()) if sp.exists() else {};assets=[]
def find_assets(obj):
 if isinstance(obj,dict):
  for k,v in obj.items():
   if isinstance(v,str) and v.endswith('.png'):
    p=Path(v);p=p if p.is_absolute() else ROOT/p
    if p.exists():assets.append((p,obj))
   elif isinstance(v,(dict,list)):find_assets(v)
 elif isinstance(obj,list):
  for v in obj:find_assets(v)
find_assets(sm)
s=Scene('11_structure_evidence','A predicted chimera is not an experimental structure','Protein architecture can be source-grounded while the three-dimensional fold remains uncertain.',
 'structures/manifest.json; PDB 6N9N; '+PAPER+' Fig. 3a / Methods',
 'Structure assets are independently supplied with sequence identity, residue-source boundaries, source coordinates, method and confidence. The paper Fig. 3a chimera is also a prediction (AlphaFold3), not an experimental structure.',
 'Gsdmd:Tmem106a model is low confidence. Experimental parent GSDMD coordinates do not validate the chimera fold.')
# Assets selected by provenance-bearing names; final build fails to claim availability if absent.
pred=next(((p,o) for p,o in assets if any(t in str(p).lower() for t in ['chimera','boltz','prediction','gsdmd_tmem','gsdmd-tmem'])),None)
exp=next(((p,o) for p,o in assets if '6n9n' in str(p).lower() or 'parent' in str(p).lower()),None)
for x,title,item,col in [(65,'Experimental parent · GSDMD',exp,B),(820,'Predicted candidate · 118 residues',pred,R)]:
 s.rect(x,269,716,439,WHITE,MID,10);s.text(x+22,285,672,45,title,28,col,True)
 if item:
  image_path=item[0].with_name('cartoon.png') if item[0].with_name('cartoon.png').exists() else item[0]
  s.image(image_path,x+24,331,666,306)
 else:s.text(x+35,388,644,153,'Structure image not available in this export.\nNo substitute structure is claimed.',30,GRAY,align='center')
 s.text(x+22,653,672,52,'PDB 6N9N · experimental mouse parent\n3.30 Å crystallographic structure' if col==B else 'Boltz2 candidate prediction · mean pLDDT 48.70\npTM 0.336 · low confidence',19,GRAY)
s.label(65,741,1470,'Blue residues 1–73: Gsdmd-derived. Red residues 74–118: novel out-of-frame C terminus.',B)

# 12 main conclusions.
s=Scene('12_next_experiment','An auditable result. An independent test comes next.','This gallery demonstrates an auditable experiment and a working evidence product.',
 'All reviewed artifacts; dataset audit; MODEL_CARD.md; compute + demo status records',
 'Summarize the actual retrospective result, independent GPU pilot, cited-report workflow, and remaining experimental needs; no newly discovered chimeras are claimed.',
 'Probe-panel selection, unknown QC, coarse Hi-C context, small labels and model uncertainty limit generalization.')
for x,title,col,items in [(66,'Demonstrated',B,['479 reproducible evaluated pairs','Gene-disjoint held-out predictions','Real GPU-generated junction output','Five reviewed source-grounded reports']),(820,'Still unestablished',R,['A consistent Hi-C ranking benefit','Candidate recovery in the GPU pilot','Assay QC / biological sample counts','A reliable experimental chimera fold'])]:
 s.rect(x,270,715,403,WHITE,MID,10);s.text(x+25,296,664,52,title,36,col,True)
 for j,t in enumerate(items):s.circle(x+30,390+j*64,12,col);s.text(x+62,374+j*64,611,58,t,26,INK)
s.rect(65,720,1470,68,PALE_B);s.text(91,738,1420,43,'Next test: complete assay/QC records + independent candidates + targeted junction confirmation.',27,B,True)

# Native PowerPoint and SVG export from exactly the same scene elements.
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9);prs.core_properties.title='chRNA evidence gallery';prs.core_properties.subject='Reviewed methods, results, candidate junctions and structure provenance';prs.core_properties.author='chRNA hackathon'
manifest=[];pdfs=[]
def rgb(c):return RGBColor.from_string(c.lstrip('#'))
for idx,s in enumerate(SLIDES,1):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);svg=[f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" width="1600" height="900" viewBox="0 0 1600 900">']
 for e in s.items:
  k=e['kind'];x=e.get('x',0);y=e.get('y',0)
  if k in ('rect','circle'):
   shape=slide.shapes.add_shape(MSO_SHAPE.OVAL if k=='circle' else MSO_SHAPE.ROUNDED_RECTANGLE if e.get('r') else MSO_SHAPE.RECTANGLE,Inches(x/100),Inches(y/100),Inches(e['w']/100),Inches(e['h']/100))
   shape._element.spPr.append(OxmlElement('a:effectLst'))
   if e.get('r') and k=='rect':shape.adjustments[0]=.08
   shape.fill.solid();shape.fill.fore_color.rgb=rgb(e['fill'])
   if e.get('stroke'):shape.line.color.rgb=rgb(e['stroke']);shape.line.width=Pt(e['sw']*.72)
   else:shape.line.fill.background()
   if k=='circle':svg.append(f'<ellipse cx="{x+e["w"]/2}" cy="{y+e["h"]/2}" rx="{e["w"]/2}" ry="{e["h"]/2}" fill="{e["fill"]}" stroke="{e.get("stroke") or "none"}" stroke-width="{e["sw"]}"/>')
   else:svg.append(f'<rect x="{x}" y="{y}" width="{e["w"]}" height="{e["h"]}" rx="{e.get("r",0)}" fill="{e["fill"]}" stroke="{e.get("stroke") or "none"}" stroke-width="{e["sw"]}"/>')
  elif k=='line':
   shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x/100),Inches(y/100),Inches(e['x2']/100),Inches(e['y2']/100));shape.line.color.rgb=rgb(e['color']);shape.line.width=Pt(e['sw']*.72);shape._element.spPr.append(OxmlElement('a:effectLst'))
   if e.get('dash'):
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    shape.line.dash_style=MSO_LINE_DASH_STYLE.DASH
   dash=' stroke-dasharray="7 7"' if e.get('dash') else ''
   svg.append(f'<line x1="{x}" y1="{y}" x2="{e["x2"]}" y2="{e["y2"]}" stroke="{e["color"]}" stroke-width="{e["sw"]}"{dash}/>')
  elif k=='text':
   shape=slide.shapes.add_textbox(Inches(x/100),Inches(y/100),Inches(e['w']/100),Inches(e['h']/100));tf=shape.text_frame;tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.vertical_anchor=MSO_ANCHOR.TOP;tf._txBody.bodyPr.set('anchorCtr','0');tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
   for j,line in enumerate(e['text'].split('\n')):
    p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(e['size']*.72);p.font.bold=e['bold'];p.font.italic=e['italic'];p.font.color.rgb=rgb(e['color']);p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(e['size']*1.18*.72);p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[e['align']]
    xx=x if e['align']=='left' else x+e['w']/2 if e['align']=='center' else x+e['w'];anchor={'left':'start','center':'middle','right':'end'}[e['align']]
    svg.append(f'<text x="{xx}" y="{y+e["size"]*.91+j*e["size"]*1.18}" font-family="DejaVu Sans, sans-serif" font-size="{e["size"]}" fill="{e["color"]}" font-weight="{700 if e["bold"] else 400}" font-style="{"italic" if e["italic"] else "normal"}" text-anchor="{anchor}">{html.escape(line)}</text>')
  elif k=='image':
   slide.shapes.add_picture(e['path'],Inches(x/100),Inches(y/100),Inches(e['w']/100),Inches(e['h']/100));b64=base64.b64encode(Path(e['path']).read_bytes()).decode();svg.append(f'<image x="{x}" y="{y}" width="{e["w"]}" height="{e["h"]}" xlink:href="data:image/png;base64,{b64}"/>')
 slide.notes_slide.notes_text_frame.text=s.notes();svg.append('</svg>');sv=''.join(svg);svgpath=G/(s.slug+'.svg');svgpath.write_text(sv)
 cairosvg.svg2png(bytestring=sv.encode(),write_to=str(G/(s.slug+'.png')),output_width=1600,output_height=900)
 pdf=G/(s.slug+'.pdf');cairosvg.svg2pdf(bytestring=sv.encode(),write_to=str(pdf));pdfs.append(pdf)
 manifest.append({'slide':idx,'slug':s.slug,'title':s.title,'png':str((G/(s.slug+'.png')).relative_to(ROOT)),'svg':str(svgpath.relative_to(ROOT)),'source':s.source,'method':s.method,'interpretation_and_caveat':s.caveat,'editable_elements':sum(e['kind']!='image' for e in s.items),'image_assets':[e['path'] for e in s.items if e['kind']=='image']})
prs.save(OUT/'chRNA_gallery.pptx')
writer=PdfWriter()
for p in pdfs:writer.append(str(p))
writer.write(OUT/'chRNA_gallery_source_render.pdf')
canvas=Image.new('RGB',(1600,4*320),(225,228,234));draw=ImageDraw.Draw(canvas)
for i,s in enumerate(SLIDES):
 im=Image.open(G/(s.slug+'.png')).convert('RGB');im.thumbnail((500,281));x=20+(i%3)*530;y=15+(i//3)*320;canvas.paste(im,(x,y));draw.text((x,y+285),f'{i+1:02}  {s.slug[3:].replace("_"," ")}',fill='black')
canvas.save(OUT/'gallery_contact_sheet.png')
inputs=['results/classifier/metrics.json','results/classifier/predictions.tsv','results/presentation/candidate_examples.json','results/presentation/reference_orf_verification.json','results/compute/pilot_summary.json','results/structures/manifest.json']
payload={'slides':manifest,'render_method':'Shared-source SVG rendering; native PPTX consists of editable text/shapes, with structure images where supplied. Actual PowerPoint/LibreOffice PDF validation tracked separately.','input_sha256':{p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in inputs if (ROOT/p).exists()},'structure_asset_count':len(assets)}
(G/'figure_manifest.json').write_text(json.dumps(payload,indent=2));(G/'scenes.json').write_text(json.dumps([{'slug':s.slug,'items':s.items} for s in SLIDES],indent=2))
print(f'Created {len(SLIDES)} editable slides; {len(assets)} discovered structure images. Output: {OUT}')
