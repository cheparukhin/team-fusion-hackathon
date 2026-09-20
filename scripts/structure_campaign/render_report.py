"""Build the source-bound, offline structure/disorder campaign atlas.

The analytical environment supplies matplotlib/pandas; the optional presentation
environment supplies python-pptx/Pillow. Missing inputs render explicit empty
states. This script never generates analytical measurements or changes inputs.
"""
from __future__ import annotations
import argparse, csv, gzip, hashlib, html, json, math, os, shutil, subprocess, tempfile
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / 'results/structure_campaign'
FIG = BASE / 'figures'
OUT = BASE / 'report'
BLUE, RED, GRAY, INK = '#0000FF', '#D22D27', '#75808C', '#172431'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none','axes.titleweight':'bold','axes.labelcolor':INK,'text.color':INK})
INPUTS = {}
FIGURES = []

def read_json(path, default=None):
    path = BASE / path
    if not path.exists(): return default if default is not None else {}
    INPUTS[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    return json.loads(path.read_text())

def read_tsv(path):
    path = BASE / path
    if not path.exists(): return []
    INPUTS[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    op = gzip.open if path.suffix == '.gz' else open
    with op(path, 'rt') as handle: return list(csv.DictReader(handle, delimiter='\t'))

def num(row, key, default=None):
    try:
        value = float(row.get(key,''))
        return value if math.isfinite(value) else default
    except (ValueError, TypeError): return default

def empty(ax, title, text):
    ax.set_axis_off(); ax.set_title(title, loc='left', pad=18)
    ax.text(.05,.62,'Not available',transform=ax.transAxes,color=GRAY,size=22,weight='bold')
    ax.text(.05,.48,text,transform=ax.transAxes,color=GRAY,size=12,wrap=True,va='top')

def figure(slug,title,source,method,interpretation):
    fig = plt.figure(figsize=(15.6,7.3), layout='constrained',facecolor='white')
    fig.suptitle(title,fontsize=23,weight='bold',x=.025,ha='left')
    FIGURES.append({'slug':slug,'title':title,'source':source,'method':method,'interpretation':interpretation})
    return fig

def save(fig, directory=None):
    directory = Path(directory) if directory is not None else FIG
    directory.mkdir(parents=True,exist_ok=True)
    item = FIGURES[-1]
    for ext in ('svg','png','pdf'):
        path = directory / f"{item['slug']}.{ext}"
        fig.savefig(path,dpi=150,bbox_inches='tight',facecolor='white')
        item[ext] = str(path.relative_to(ROOT))
    if fig._suptitle is not None: fig._suptitle.set_visible(False)
    slide_path=directory/f"{item['slug']}.slide.png";fig.savefig(slide_path,dpi=150,bbox_inches='tight',facecolor='white');item['slide_png']=str(slide_path.relative_to(ROOT))
    plt.close(fig)

def clean_tier(value):
    return str(value or 'Unspecified tier').replace('_',' ')

def protocol_label(value):
    if value.endswith('_precomputed'): return 'Cached MSA-backed'
    if value.endswith('_single_sequence'): return 'Single-sequence'
    return value.replace('boltz2_2.2.1_','').replace('_',' ')

def render_cross_model_structures(snapshot):
    """Render verified single-chain coordinates; never alter the source contract."""
    import gemmi
    from pymol import cmd
    snapshot=json.loads(json.dumps(snapshot));target=BASE/'cross_model_gsdmd/report/structures';target.mkdir(parents=True,exist_ok=True)
    expected=snapshot['sequence_sha256'];receipts=[]
    for model in snapshot.get('models',[]):
        if model.get('status') not in {'verified','cached_verified','completed'}: continue
        assert model.get('sequence_verified') is True
        path=ROOT/model['model_path'];digest=hashlib.sha256(path.read_bytes()).hexdigest()
        model_id=model['model_id'];assert all(c.isalnum() or c in '_-.' for c in model_id)
        dest=target/(model_id+'.png');receipt_path=target/(model_id+'.json')
        cached=json.loads(receipt_path.read_text()) if receipt_path.exists() else {}
        if not (cached.get('render_settings_version')==2 and cached.get('model_sha256')==digest and cached.get('sequence_sha256')==expected and dest.exists() and cached.get('png_sha256')==hashlib.sha256(dest.read_bytes()).hexdigest()):
            structure=gemmi.read_structure(str(path));assert len(structure)==1
            matches=[]
            for chain in structure[0]:
                residues=[r for r in chain if 'CA' in r]
                sequence=''.join(gemmi.find_tabulated_residue(r.name).one_letter_code for r in residues)
                if hashlib.sha256(sequence.encode()).hexdigest()==expected: matches.append((chain.name,residues,sequence))
            assert len(matches)==1, 'Require one complete exact-sequence chain'
            chain,residues,sequence=matches[0];assert len(sequence)==118
            assert [r.seqid.num for r in residues]==list(range(1,119)), 'One-based uninterrupted residue numbering required'
            assert all(not r.seqid.icode.strip() for r in residues), 'Insertion-code mapping requires explicit handling'
            cmd.reinitialize();cmd.load(str(path),'hypothesis');cmd.remove('not (polymer.protein and chain '+chain+')');cmd.hide('everything');cmd.show('cartoon')
            cmd.set_color('source_blue',[0,0,1]);cmd.set_color('source_red',[210/255,45/255,39/255]);cmd.color('source_red');cmd.color('source_blue','resi 1-73')
            cmd.set('cartoon_loop_radius',.25);cmd.set('cartoon_fancy_helices',1);cmd.set('cartoon_fancy_sheets',1);cmd.set('cartoon_sampling',14);cmd.set('ray_opaque_background',0);cmd.set('ray_shadows',0);cmd.set('orthoscopic',1);cmd.set('antialias',2);cmd.bg_color('white');cmd.viewport(1800,1400);cmd.orient();cmd.turn('y',15);cmd.zoom('all',buffer=5,complete=1)
            cmd.png(str(dest),width=1800,height=1400,dpi=250,ray=1)
            cached={'render_settings_version':2,'pymol_version':cmd.get_version()[0],'model_id':model_id,'model_path':str(path.relative_to(ROOT)),'model_sha256':digest,'sequence_sha256':expected,'sequence_verified_from_coordinate_residues':True,'length_aa':118,'chain':chain,'residue_numbers':list(range(1,119)),'colors':{'blue_parent_A_RNA':'1-73','red_parent_B_or_split_RNA':'74-118','split_junction_codon':74},'view':list(cmd.get_view()),'renderer':'PyMOL coordinate cartoon; independent orientation per model, not structural alignment','png_sha256':hashlib.sha256(dest.read_bytes()).hexdigest()}
            receipt_path.write_text(json.dumps(cached,indent=2))
        model['png']=str(dest.relative_to(ROOT));model['render_receipt']=str(receipt_path.relative_to(ROOT));receipts.append(cached)
    (target/'manifest.json').write_text(json.dumps({'sequence_sha256':expected,'models':receipts},indent=2))
    return snapshot

def cross_model_comparison(snapshot=None, draft=False):
    """Render only supplied, audited records; an absent contract changes nothing."""
    if snapshot is None:
        snapshot=read_json('cross_model_gsdmd/comparison.json')
    if not snapshot:
        return
    expected='f0766d124b52f0061597ce4e822e9275a04152df574f9a512631a0fa6ed8a2fa'
    assert snapshot.get('sequence_sha256')==expected and snapshot.get('length_aa')==118
    if not draft: snapshot=render_cross_model_structures(snapshot)
    for relative,expected_hash in snapshot.get('inputs_sha256',snapshot.get('input_sha256',{})).items():
        path=ROOT/relative;actual_hash=hashlib.sha256(path.read_bytes()).hexdigest();assert actual_hash==expected_hash,relative
        INPUTS[relative]=actual_hash
    models=snapshot.get('models',[])
    for model in models:
        if not model.get('display_label'):
            if model.get('engine_family')=='Boltz': model['display_label']='Boltz2 MSA (cached)' if model.get('msa_mode')=='precomputed' else 'Boltz2 SS · seed '+str(model.get('seed','?'))[-2:]
            else: model['display_label']=model.get('engine',model['model_id'])
    by_id={m['model_id']:m for m in models}
    for model in models:
        if model.get('status') in {'verified','cached_verified','completed'}:
            assert model.get('sequence_sha256')==expected and model.get('length_aa')==118
            assert model.get('sequence_verified') is True, 'Cross-model sequence audit required'
            for field in ('model_path','png','residue_confidence_tsv','render_receipt'):
                if model.get(field):
                    path=ROOT/model[field];assert path.is_file(),str(path)
                    INPUTS[str(path.relative_to(ROOT))]=hashlib.sha256(path.read_bytes()).hexdigest()
            mean=num(model,'mean_plddt')
            if mean is not None: assert 0<=mean<=100
    groups=snapshot.get('display_groups',[])
    if not groups:
        seen=set()
        for model in models:
            key=(model.get('engine'),model.get('protocol_id'))
            if key not in seen:
                engine=model.get('engine','Unspecified engine');msa=model.get('msa_mode','');complete=model.get('status') in {'verified','cached_verified','completed'}
                if model.get('engine_family')=='Boltz': group_title='Boltz2 · cached MSA-backed' if 'precomputed' in msa else 'Boltz2 · single-sequence baseline'
                elif model.get('engine_family')=='AlphaFold2': group_title='AlphaFold2 · cached MSA\n'+('New prediction' if complete else model.get('status','Pending').replace('_',' ').capitalize())
                elif model.get('engine_family')=='ESMFold': group_title='ESMFold · single sequence\n'+('New prediction' if complete else model.get('status','Pending').replace('_',' ').capitalize())
                else: group_title=engine+(' · '+msa.replace('_',' ') if msa else '')
                groups.append({'title':model.get('display_group',group_title),'representative_model_id':model['model_id']})
                seen.add(key)
    assert len(groups)<=4, 'Choose at most four prespecified protocol representatives for this slide'
    title='Gsdmd–Tmem106a: predictions disagree on the retained segment'
    if draft: title='DRAFT · cross-model layout; new engine results pending'
    f=figure('DRAFT_cross_model_layout' if draft else '10_cross_model_gsdmd',title,
        snapshot.get('source_description','cross_model_gsdmd/comparison.json; exact-sequence model audits, coordinates and confidence records'),
        'Compare the same reference-reconstructed 118-aa peptide. Representatives follow the prespecified source order, never a new confidence-based selection. Source colors: parent-A RNA residues 1–73 blue; split codon 74 and parent-B RNA 75–118 red. Cartoons use independent orientations and display scales; they are not structural superpositions. Engine confidence is not assumed calibrated across methods.',
        'Computational validation checks prediction sensitivity, not an experimental fusion fold. Low pLDDT is not a disorder fraction; sequence-disorder predictions and model confidence remain separate.')
    grid=f.add_gridspec(2,3,width_ratios=[1,1,1.12]);cards=[f.add_subplot(grid[r,c]) for r,c in [(0,0),(0,1),(1,0),(1,1)]]
    for ax,group in zip(cards,groups):
        model=by_id[group['representative_model_id']];ax.axis('off')
        ax.set_title(group['title'],fontsize=12)
        image_path=ROOT/model['png'] if model.get('png') else None
        if image_path and image_path.is_file() and model.get('status') in {'verified','cached_verified','completed'}:
            picture=plt.imread(image_path)
            # Remove transparent canvas padding only; retain every rendered atom/ribbon pixel.
            if picture.ndim==3 and picture.shape[2]==4:
                visible=np.argwhere(picture[:,:,3]>0)
                if len(visible):
                    y0,x0=visible.min(axis=0);y1,x1=visible.max(axis=0)+1;pad=30
                    picture=picture[max(0,y0-pad):min(picture.shape[0],y1+pad),max(0,x0-pad):min(picture.shape[1],x1+pad)]
            ax.imshow(picture)
            label=model.get('display_label',model['model_id'])
            ax.text(.02,-.02,label,transform=ax.transAxes,fontsize=8,color=GRAY)
        else:
            ax.text(.5,.57,'Unavailable / pending',ha='center',transform=ax.transAxes,fontsize=19,color=GRAY,weight='bold')
            ax.text(.5,.40,model.get('unavailable_reason',model.get('reason','Status: '+model.get('status','unknown').replace('_',' ')+'. No audited result supplied.')),ha='center',va='top',transform=ax.transAxes,fontsize=10,color=GRAY,wrap=True)
    for ax in cards[len(groups):]:ax.axis('off')
    right=f.add_subplot(grid[:,2]);right.axis('off');right.set_title('Exact sequence · 118 aa',loc='left',fontsize=15)
    right.add_patch(plt.Rectangle((.02,.94),.92*73/118,.025,color=BLUE,transform=right.transAxes));right.add_patch(plt.Rectangle((.02+.92*73/118,.94),.92*45/118,.025,color=RED,transform=right.transAxes))
    right.text(.02,.90,'1–73: parent A     74: split codon     75–118: parent B',fontsize=8,transform=right.transAxes)
    right.text(.02,.83,'Model confidence — separate by protocol',fontsize=11,weight='bold',transform=right.transAxes)
    rows=[]
    for model in models:
        mean=num(model,'mean_plddt');low=num(model,'fraction_plddt_below50')
        rows.append([model.get('display_label',model['model_id']),f'{mean:.1f}' if mean is not None else 'Unavailable',f'{low:.1%}' if low is not None else '—'])
    table=right.table(cellText=rows,colLabels=['Model / seed','Mean pLDDT','pLDDT <50'],colWidths=[.55,.24,.21],bbox=[.01,.45,.97,.34],cellLoc='left');table.auto_set_font_size(False);table.set_fontsize(8)
    for (r,c),cell in table.get_celld().items():
        cell.set_edgecolor('#DCE0E7')
        if r==0:cell.set_facecolor('#EDF0F8');cell.set_text_props(weight='bold')
    right.text(.02,.39,'Sequence disorder · f_IDR by source region',fontsize=11,weight='bold',transform=right.transAxes)
    disorder_rows=[]
    for predictor in snapshot.get('disorder_methods',snapshot.get('disorder_predictors',[])):
        assert predictor.get('sequence_sha256')==expected
        fraction=num(predictor,'fraction_disordered',num(predictor,'f_idr'));assert fraction is not None and 0<=fraction<=1
        label=predictor.get('method',predictor.get('name',predictor.get('label','Predictor')))
        donor=num(predictor,'donor_fraction_disordered');tail=num(predictor,'tail_fraction_disordered')
        disorder_rows.append([label,f'{fraction:.1%}',f'{donor:.1%}' if donor is not None else '—',f'{tail:.1%}' if tail is not None else '—'])
    disorder_table=right.table(cellText=disorder_rows,colLabels=['Predictor','Whole','Donor\n1–73','Tail\n74–118'],colWidths=[.43,.19,.19,.19],bbox=[.01,.16,.97,.20],cellLoc='left');disorder_table.auto_set_font_size(False);disorder_table.set_fontsize(8.5)
    for (r,c),cell in disorder_table.get_celld().items():
        cell.set_edgecolor('#DCE0E7')
        if r==0:cell.set_facecolor('#EDF0F8');cell.set_text_props(weight='bold')
    right.text(.02,.105,'Computational validation / sensitivity only.\nConfidence ≠ disorder ≠ function.\nNo experimental fusion structure ground truth.',fontsize=9,color=RED,transform=right.transAxes,va='top',wrap=True)
    actual_models=[m for m in models if m.get('status') in {'verified','cached_verified','completed'}]
    FIGURES[-1]['comparison_scope']={'length_aa':118,'sequence_sha256':expected,'verified_model_records':len(actual_models),'new_model_records':sum(not m.get('reused_prior_prediction',False) for m in actual_models),'folding_families':sorted({m.get('engine_family',m.get('engine','Unspecified')) for m in actual_models}),'main_campaign_denominators_unchanged':True}
    target=BASE/'cross_model_gsdmd/report';save(f,target)
    (target/('DRAFT_comparison_snapshot.json' if draft else 'comparison_snapshot.json')).write_text(json.dumps(snapshot,indent=2))
    if not draft:
        links=[]
        for model in models:
            artifact_links=[]
            for field,label in [('model_path','Coordinates'),('residue_confidence_tsv','Residue confidence')]:
                if model.get(field) and (ROOT/model[field]).is_file(): artifact_links.append('<a href="'+html.escape(os.path.relpath(ROOT/model[field],target))+'">'+label+'</a>')
            links.append('<li><b>'+html.escape(model.get('display_label',model['model_id']))+'</b> — '+html.escape(model.get('status','unspecified'))+'; '+' · '.join(artifact_links)+'</li>')
        record=FIGURES[-1]
        page='<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Gsdmd cross-model comparison</title><style>body{font:17px/1.5 system-ui;max-width:1500px;margin:auto;padding:24px;color:#172431}img{width:100%}a{color:#0000ff}.note{border-left:5px solid #d22d27;padding:12px;background:#fceeed}</style></head><body><h1>'+html.escape(title)+'</h1><p class="note">'+html.escape(record['interpretation'])+'</p><img src="10_cross_model_gsdmd.svg" alt="Exact-sequence structure, confidence and disorder comparison"><p>'+html.escape(record['method'])+'</p><h2>Actual model artifacts and method status</h2><ul>'+''.join(links)+'</ul><p><a href="comparison_snapshot.json">Audited comparison snapshot</a> · <a href="../../report/structure_campaign_gallery.pptx">Campaign PowerPoint</a> · <a href="10_cross_model_gsdmd.pdf">Comparison PDF</a></p></body></html>'
        unavailable='<h2>Other audited methods</h2><ul>'+''.join('<li><b>'+html.escape(m.get('engine','Method'))+'</b> — '+html.escape(m.get('status','unknown').replace('_',' '))+': '+html.escape(m.get('reason','No successful prediction supplied.'))+'</li>' for m in snapshot.get('unavailable_methods',[]))+'</ul>'
        predictor_notes='<h2>Disorder methods and interpretation</h2><ul>'+''.join('<li>'+html.escape(m.get('method',m.get('name',m.get('label','Predictor'))))+': '+html.escape(m.get('threshold','Threshold recorded in source snapshot'))+'. '+html.escape(m.get('independence_note',''))+'</li>' for m in snapshot.get('disorder_methods',snapshot.get('disorder_predictors',[])))+'</ul>'
        page=page.replace('</body>',unavailable+predictor_notes+'</body>')
        (target/'index.html').write_text(page)
        record['related_report']=str((target/'index.html').relative_to(ROOT))
    (target/('DRAFT_figure_record.json' if draft else 'figure_record.json')).write_text(json.dumps(FIGURES[-1],indent=2))
    if not draft: (target/'render_manifest.json').write_text(json.dumps({'status':'available_source_snapshot_rendered','renderer_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'input_sha256':dict(INPUTS),'figure':FIGURES[-1]},indent=2))

def render():
    FIG.mkdir(parents=True,exist_ok=True); OUT.mkdir(parents=True,exist_ok=True)
    paper=ROOT/'docs/Venezia_et_al_2026_functional_chimeric_mRNAs.pdf'
    if paper.exists(): INPUTS[str(paper.relative_to(ROOT))]=hashlib.sha256(paper.read_bytes()).hexdigest()
    summary = read_json('analysis/summary.json')
    alternative_summary=read_json('analysis/alternative_start/summary.json')
    disorder = read_tsv('analysis/disorder_summary.tsv')
    residues = read_tsv('analysis/residue_disorder.tsv.gz')
    metrics = read_tsv('compute/model_metrics.tsv')
    all_metrics=metrics
    # A repeated diagnostic seed is not an extra sampled peptide. Display the
    # lowest recorded seed once per exact sequence and role; retain raw job data.
    unique_models={}
    for row in sorted(metrics,key=lambda r:(num(r,'seed',float('inf')),r.get('job_id',''))):
        key=(row.get('sequence_sha256') or row.get('peptide_id'),row.get('role','candidate'),row.get('protocol_id','unspecified'))
        if key not in unique_models: unique_models[key]=row
    metrics=[r for r in unique_models.values() if num(r,'mean_plddt') is not None]
    jobs = read_json('compute/jobs.json')
    compute_summary = read_json('compute/summary.json')
    compute_receipt = read_json('compute/campaign_receipt.json')
    cleanup_receipt = read_json('compute/cleanup_receipt.json')
    audit = read_json('cohort/manifest.json')
    peptides = read_tsv('cohort/peptides.tsv')
    by_peptide_id={r['peptide_id']:r for r in peptides}
    for row in disorder:
        if row.get('role','candidate')=='candidate' and row.get('peptide_id') in by_peptide_id:
            source=by_peptide_id[row['peptide_id']]
            assert row['sequence_sha256']==source['sequence_sha256'], 'Disorder/sequence hash mismatch'
            assert num(row,'length')==num(source,'length_aa'), 'Disorder/sequence length mismatch'
        if num(row,'f_idr') is not None: assert 0<=num(row,'f_idr')<=1
    source_regions = read_tsv('cohort/regions.tsv')
    orfs = read_tsv('cohort/orf_hypotheses.tsv')
    eligibility = read_tsv('cohort/eligibility_funnel.tsv')
    diversity = read_json('analysis/diversity_summary.json')
    diversity_models=read_tsv('diversity/model_audit.tsv')
    diversity_domains=read_tsv('diversity/domain_audit.tsv')
    rarefaction=read_tsv('diversity/rarefaction.tsv')
    gate_sensitivity=read_tsv('diversity/gate_sensitivity.tsv')
    architectures=read_tsv('domains/architectures.tsv')
    domain_hits=read_tsv('domains/domain_hits.tsv')
    domain_manifest=read_json('domains/manifest.json')
    structure_gallery = read_json('compute/structure_gallery.json')
    structure_assets = structure_gallery if isinstance(structure_gallery,list) else structure_gallery.get('assets',[])
    structure_assets=sorted(structure_assets,key=lambda a:not(a.get('title')=='Gsdmd:Tmem106a' and 'precomputed' in a.get('protocol_id','')))
    for asset in structure_assets:
        if asset.get('png') and (ROOT/asset['png']).exists(): INPUTS[asset['png']]=hashlib.sha256((ROOT/asset['png']).read_bytes()).hexdigest()
    regions = read_tsv('analysis/region_metrics.tsv')
    region_peptides=read_tsv('analysis/region_peptide_summary.tsv')
    paired_fragments=read_tsv('analysis/paired_controls.tsv')
    paired_native=read_tsv('analysis/paired_native_regions.tsv')
    region_summary=read_json('analysis/region_summary.json')
    census_summary=read_json('analysis/census_summary.json')
    novel_tail_summary=read_json('analysis/novel_tail_summary.json')
    # Retain roles; parent controls never silently enter peptide-hypothesis summaries.
    hypotheses = [r for r in disorder if r.get('role','candidate') not in ('parent','control','native_control','matched_native','parent_control')]
    tiers = sorted({r.get('cohort_tier','unspecified') for r in hypotheses})
    colors = [BLUE, RED, '#7156AA', '#C7862A', GRAY]
    f = figure('01_eligibility_funnel','RNA evidence is not a protein denominator',
        'cohort audit and eligibility tables',
        'Keep RNA-pair, exact-junction, coding-hypothesis and unique-peptide counts distinct. Strict and conditional tiers are not combined.',
        'Local probe support does not validate a full transcript or translated protein. An empty strict-primary set is a result, not a zero-disorder estimate.')
    ax, bx = f.subplots(1,2,gridspec_kw={'width_ratios':[1.15,1]})
    if audit.get('counts'):
        counts=audit['counts'];ax.axis('off')
        entries=[('RNA reference pairs',counts.get('primary_reference_pairs')),
                 ('Pairs with annotated-start hypotheses',counts.get('conditional_census_pairs')),
                 ('Conditional annotated-start peptides',counts.get('conditional_census_unique_peptides')),
                 ('Strict-primary eligible peptides',counts.get('strict_primary_peptides'))]
        for i,(label,value) in enumerate(entries):
            y=.92-i*.23;ax.text(.02,y,str(value) if value is not None else 'Unavailable',size=35,weight='bold',color=RED if i==3 else BLUE,transform=ax.transAxes);ax.text(.25,y+.01,label,size=13,transform=ax.transAxes,va='center')
        ax.set_title('Evidence and reconstruction units',loc='left')
        ax.text(.02,-.03,f"Conditional annotated-start census: {counts.get('conditional_census_unique_peptides','?')} peptides / {counts.get('conditional_census_pairs','?')} pairs.\nFull-chain evidence remains unresolved.",size=12,color=RED,transform=ax.transAxes)
    elif eligibility:
        labels=[r.get('stage',r.get('label','')) for r in eligibility]
        values=[num(r,'count',0) for r in eligibility]
        ax.barh(range(len(values)),values,color=[BLUE if i%2==0 else RED for i in range(len(values))]);ax.set_yticks(range(len(values)),labels);ax.invert_yaxis();ax.set_xlabel('Recorded count; units change at each stage')
        for i,v in enumerate(values): ax.text(v,i,f' {v:,.0f}',va='center')
    else: empty(ax,'Eligibility audit','Actual RNA → junction → complete ORF → peptide counts await the audited cohort table.')
    if audit.get('counts'):
        c=audit['counts'];cons=c['conditional_consensus_unique_peptides'];total=c['conditional_census_unique_peptides'];alternative=c['alternative_start_only_unique_peptides']
        bx.barh(0,total-cons,color=BLUE,label='Other conditional annotated-start');bx.barh(0,cons,left=total-cons,color=RED,label='Consensus subset');bx.barh(1,alternative,color=GRAY,label='Alternative-start sensitivity')
        bx.set_yticks([0,1],['Conditional census','Separate alternative starts']);bx.invert_yaxis();bx.set_xlabel('Unique reference-derived peptide hypotheses');bx.set_title('188 main conditional peptides; 30 nested consensus')
        bx.text(total,0,f' {total} ({cons} consensus)',va='center');bx.text(alternative,1,f' {alternative}',va='center');bx.legend(loc='lower right',fontsize=9)
        bx.text(0,-.12,f"All-start reconstruction inventory: {c['all_unique_peptide_hypotheses']} unique peptides; not the folded denominator.",transform=bx.transAxes,fontsize=10,color=GRAY)
    else: empty(bx,'Sequence inventory','No peptide census has been supplied. The 109-pair RNA reference is not a count of proteins.')
    save(f)

    f=figure('02_disorder_census','Sequence disorder: census and predictor sensitivity',
        'analysis/disorder_summary.tsv; analysis/summary.json',
        'f_IDR is the called-disordered residue fraction: V3 score ≥0.5; V1 score ≥0.42. Predominantly disordered means protein f_IDR >0.5; sensitivity cutoffs are 0.4 and 0.6. No finite-census bootstrap intervals.',
        'A metapredict version comparison is within-method sensitivity, not independent agreement. Retained-parent comparisons are separate; native population matching is unavailable.')
    axs=f.subplots(1,3)
    conditional=[r for r in hypotheses if r.get('cohort_tier') in ('conditional_consensus','conditional_annotated_start')]
    consensus=[r for r in conditional if r.get('cohort_tier')=='conditional_consensus']
    groups=[('Conditional census',conditional,BLUE),('Consensus subset',consensus,RED)]
    if conditional:
        for label,group,color in groups:
            vals=sorted(num(r,'f_idr') for r in group if num(r,'f_idr') is not None)
            if vals:
                axs[0].step(vals,np.arange(1,len(vals)+1)/len(vals),where='post',color=color,label=f'{label} (n={len(vals)})')
                axs[1].plot([.4,.5,.6],[np.mean(np.array(vals)>cut) for cut in [.4,.5,.6]],'o-',color=color,label=label)
        axs[0].axvline(.5,color=GRAY,ls=':');axs[0].set(xlim=(0,1),ylim=(0,1.02),xlabel='V3 called-disordered residue fraction',ylabel='Fraction of peptide hypotheses',title='Census and nested consensus subset');axs[0].legend(fontsize=8,loc='lower right')
        axs[1].set(xlabel='Predominantly-disordered cutoff (strict >)',ylabel='Fraction exceeding cutoff',ylim=(0,1),xticks=[.4,.5,.6],title='Threshold sensitivity')
        n3=sum(num(r,'f_idr',0)>.5 for r in conditional);n1=sum(num(r,'v1_f_idr',0)>.5 for r in conditional)
        axs[1].text(.5,.9,f'V3: {n3}/{len(conditional)} = {n3/len(conditional):.1%}\nV1: {n1}/{len(conditional)} = {n1/len(conditional):.1%}',transform=axs[1].transAxes,ha='center',va='top',fontsize=14,color=INK,weight='bold')
    else:
        empty(axs[0],'Disorder distribution','metapredict disorder outputs have not arrived.');empty(axs[1],'Threshold sensitivity','No computed fractions available.')
    sensitivity=[r for r in conditional if num(r,'v1_f_idr') is not None and num(r,'f_idr') is not None]
    if sensitivity:
        for label,group,color in groups:
            group=[r for r in group if r in sensitivity]
            axs[2].scatter([num(r,'f_idr') for r in group],[num(r,'v1_f_idr') for r in group],color=color,s=20,alpha=.6,label=label)
        axs[2].plot([0,1],[0,1],color=GRAY,ls=':');axs[2].set(xlim=(0,1),ylim=(0,1),xlabel='metapredict V3 f_IDR',ylabel='metapredict V1 f_IDR',title='Within-method sensitivity')
        axs[2].text(0,-.24,'Related networks; not independent validation.\nIUPred unavailable; retained-parent comparisons in appendix.',transform=axs[2].transAxes,size=9,color=GRAY)
    else: empty(axs[2],'Independent method / matched native','No independent predictor or matched-native comparison is inferred from one-method results. Availability is reported in the analysis methods.')
    save(f)

    f=figure('03_disorder_confidence','Predicted disorder and fold confidence answer different questions',
        'analysis/disorder_summary.tsv; compute/model_metrics.tsv',
        'Join exact sequence hashes. Each technically complete first-seed model is shown within its protocol; low confidence is retained. Single-sequence mode lacks evolutionary MSA context and is interpreted separately.',
        'pLDDT is confidence, not disorder or function. V3 training includes AlphaFold-derived information; agreement is not independent validation. Named examples are illustrative; short peptides require particular caution.')
    by_hash={r.get('sequence_sha256'):r for r in hypotheses if r.get('sequence_sha256')};by_id={r.get('peptide_id'):r for r in hypotheses};joined=[]
    for m in metrics:
        if m.get('role','candidate') in ('parent','control','native_control','matched_native','parent_control'): continue
        d=by_hash.get(m.get('sequence_sha256')) or by_id.get(m.get('peptide_id'))
        if d and num(d,'f_idr') is not None and num(m,'mean_plddt') is not None:
            assert d.get('sequence_sha256')==m.get('sequence_sha256'), 'Model/disorder sequence hash mismatch'
            assert 0<=num(m,'mean_plddt')<=100
            joined.append((d,m))
    protocols=sorted({m.get('protocol_id','unspecified') for d,m in joined})
    axes=np.atleast_1d(f.subplots(1,max(1,len(protocols))+1));bx=axes[-1]
    if joined:
        for ax,protocol in zip(axes,protocols):
            protocol_rows=[(d,m) for d,m in joined if m.get('protocol_id','unspecified')==protocol]
            for i,tier in enumerate(tiers):
                pairs=[(d,m) for d,m in protocol_rows if d.get('cohort_tier','unspecified')==tier]
                if pairs: ax.scatter([num(d,'f_idr') for d,m in pairs],[num(m,'mean_plddt') for d,m in pairs],s=60,color=colors[i%len(colors)],label=clean_tier(tier),alpha=.8)
            named_examples={'Zc3h7a:Ppp4r1','Psmc2:Pmpcb','Gsdmd:Tmem106a'}
            labels=[(d,m) for d,m in protocol_rows if len(protocol_rows)<=8 or d.get('pair_ids') in named_examples]
            for d,m in labels:
                label=d.get('pair_ids','')[:28]+f" ({int(num(d,'length',0))} aa)"
                ax.annotate(label,(num(d,'f_idr'),num(m,'mean_plddt')),xytext=(-8,8) if num(d,'f_idr')>.6 else (8,8),ha='right' if num(d,'f_idr')>.6 else 'left',textcoords='offset points',fontsize=7.5)
            ax.axvline(.5,color=GRAY,ls=':');ax.axhline(70,color=GRAY,ls=':');ax.set(xlim=(0,1.03),ylim=(0,100),xlabel='Sequence-predicted f_IDR',ylabel='Mean model pLDDT (0–100)',title=protocol_label(protocol)+f' (n={len(protocol_rows)})');ax.legend(fontsize=8,loc='lower left')
        bx.axis('off');bx.text(.05,.86,f'{len(joined)} model records\n(candidate × protocol)\n{len(hypotheses)} sequence predictions',size=18,weight='bold');ledger=compute_summary.get('protocols',{})
        ledger_text='\n'.join(protocol_label(k)+f": {v['verified_models']}/{v['planned_models']} verified" for k,v in ledger.items()) if ledger else f"{compute_summary.get('verified_models','?')}/{compute_summary.get('planned_models','?')} jobs verified"
        bx.text(.05,.61,ledger_text+'\n(controls and extra seed jobs included)',size=13);bx.text(.05,.34,'Separate protocol panels.\nNo pooled confidence estimate.\nSelected panel, not prevalence.\nMissing structures remain missing.',size=14)
    else:
        empty(axes[0],'Sequence-matched scatter','No sequence-matched confidence measurements available. No synthetic points are drawn.');empty(bx,'Structural coverage','Technical completion and low confidence are separate outcomes. Sequence-only results remain usable.')
    save(f)

    f=figure('04_junction_profiles','Junction-centered profiles and source-colored structures',
        'Venezia et al., Nature 2026, Fig. 3 (doi:10.1038/s41586-026-10982-x); analysis/residue_disorder.tsv.gz; audited residue map; compute model artifacts',
        'Residue profiles retain one-based positions. Blue/red provenance requires a verified residue-source map, including split junction codons; no boundary is guessed.',
        'Gsdmd:Tmem106a is a paper functional exemplar: predicted disorder or low confidence does not exclude function. Predictions are not physical disorder measurements.')
    ax,bx=f.subplots(1,2,gridspec_kw={'width_ratios':[1.3,1]})
    if residues:
        present_ids={r.get('peptide_id') for r in residues}
        preferred=next((a.get('peptide_id') for a in structure_assets if a.get('peptide_id') in present_ids),None)
        first=preferred or next((r.get('peptide_id') for r in hypotheses if r.get('peptide_id') in present_ids),residues[0].get('peptide_id'))
        profile=[r for r in residues if r.get('peptide_id')==first and num(r,'score') is not None]
        label=next((r.get('pair_ids',first) for r in hypotheses if r.get('peptide_id')==first),first)
        ax.plot([num(r,'pos1') for r in profile],[num(r,'score') for r in profile],color=BLUE,lw=1.5);ax.set(xlabel='Residue position (one-based)',ylabel='metapredict disorder score',title=f'Inventory example: {label[:55]}',ylim=(0,1))
        ax.axhline(.5,color=GRAY,lw=1,ls=':')
        example=next((a for a in structure_assets if a.get('peptide_id')==first),{})
        boundary=example.get('junction_residue_1based')
        if boundary is None: boundary=next((x['start_residue_1based'] for x in example.get('source_regions',[]) if x.get('source_class')=='junction_split_codon'),None)
        for region in example.get('source_regions',[]):
            ax.axvspan(region['start_residue_1based']-.5,region['end_residue_1based']+.5,color=BLUE if region.get('source_class','').startswith('parent_a') else RED,alpha=.07,zorder=0)
        if boundary is not None:
            ax.axvline(float(boundary),color=RED,ls='--',lw=1);ax.text(float(boundary),.98,' junction',color=RED,va='top')
        confidence_path=example.get('residue_confidence_tsv')
        if confidence_path:
            cp=ROOT/confidence_path
            if cp.exists():
                INPUTS[str(cp.relative_to(ROOT))]=hashlib.sha256(cp.read_bytes()).hexdigest()
                cf=list(csv.DictReader(cp.open(),delimiter='\t'));cx=[num(r,'position',num(r,'pos1',num(r,'residue_1based',num(r,'residue')))) for r in cf];cy=[num(r,'plddt') for r in cf]
                if all(x is not None for x in cx+cy):
                    right=ax.twinx();right.plot(cx,cy,color=GRAY,lw=1,alpha=.7);right.set_ylim(0,100);right.set_ylabel('pLDDT (0–100), gray',color=GRAY)
        record=next((r for r in hypotheses if r.get('peptide_id')==first),{})
        detail=f"V3 f_IDR {num(record,'f_idr',float('nan')):.3f}; related V1 {num(record,'v1_f_idr',float('nan')):.3f}. Example, not a prevalence sample."
        ax.text(0,-.22,detail,transform=ax.transAxes,size=10,wrap=True)
    else: empty(ax,'Residue profiles','Actual per-residue disorder and confidence profiles have not arrived.')
    visible_asset=next((a for a in structure_assets if a.get('png') and (ROOT/a['png']).exists() and any(r.get('peptide_id')==a.get('peptide_id') and r.get('sequence_sha256')==a.get('sequence_sha256') for r in hypotheses)),None)
    if visible_asset:
        asset_path=ROOT/visible_asset['png'];INPUTS[str(asset_path.relative_to(ROOT))]=hashlib.sha256(asset_path.read_bytes()).hexdigest();bx.imshow(plt.imread(asset_path));bx.axis('off');bx.set_title(visible_asset.get('title','Sequence-matched predicted structure')+' · cached MSA-backed prediction',fontsize=12);bx.text(.02,-.05,f"Mean pLDDT {visible_asset.get('mean_plddt',float('nan')):.1f} · {visible_asset.get('confidence_status','')}\nBlue: RNA parent A; red: parent B / split codon.",transform=bx.transAxes,size=10,color=GRAY)
    else: empty(bx,'Blue/red structures','Campaign models and verified residue-source boundaries are pending. Existing reference models are not substituted into this cohort.')
    save(f)

    f=figure('05_diversity_coverage','Sequence families and confidence coverage are different evidence',
        'domains/architectures.tsv; domain_hits.tsv; diversity/model_audit.tsv; domain_audit.tsv; gate_sensitivity.tsv',
        'Pfam 38.2 model-specific gathering thresholds on the 188 conditional sequences. Family frequency counts each peptide once per family; all passing overlapping hits retained. Model confidence is separated by protocol.',
        'Sequence families are not 3D clusters; no hit is not fold novelty. Overlapping spans are not independent domains. Coverage depends on the fixed confidence gate.')
    sensitivity_rows=[r for r in gate_sensitivity if r.get('population')=='candidate' and r.get('protocol_id','').endswith('single_sequence')]
    if sensitivity_rows:
        details='; '.join(f"pLDDT ≥{int(num(r,'residue_plddt_threshold'))}: {int(num(r,'peptides_with_passing_spans'))}/{int(num(r,'actual_validated_peptide_denominator'))}" for r in sensitivity_rows)
        FIGURES[-1]['interpretation']+=' Candidate gate sensitivity ('+details+'); ≥50-aa span and ≥80% passing residues fixed. No alternate clustering claimed.'
    ax,bx,cx=f.subplots(1,3,gridspec_kw={'width_ratios':[.85,1.15,1]})
    candidate_arch=[r for r in architectures if r.get('role')=='conditional_candidate']
    candidate_hits=[r for r in domain_hits if r.get('role')=='conditional_candidate']
    if candidate_arch:
        n_hit=sum(num(r,'pfam_domain_hit_count',0)>0 for r in candidate_arch);n_no=len(candidate_arch)-n_hit
        ax.barh([0,1],[n_hit,n_no],color=[BLUE,GRAY]);ax.set_yticks([0,1],['Pfam GA hit','No Pfam GA hit']);ax.invert_yaxis();ax.set_xlabel('Conditional peptide hypotheses');ax.set_title('Sequence annotation coverage')
        for i,v in enumerate([n_hit,n_no]):ax.text(v,i,f' {v}/{len(candidate_arch)}',va='center')
        family_members={};family_names={}
        for r in candidate_hits:
            family_members.setdefault(r['pfam_accession'],set()).add(r['sequence_id']);family_names[r['pfam_accession']]=r['pfam_name']
        top=sorted(family_members,key=lambda k:(-len(family_members[k]),k))[:9]
        bx.barh(range(len(top)),[len(family_members[k]) for k in top],color=BLUE);bx.set_yticks(range(len(top)),[family_names[k] for k in top]);bx.invert_yaxis();bx.set_xlabel('Peptides with family (nonexclusive)');bx.set_title(f'{len(family_members)} detected Pfam families')
    else:empty(ax,'Sequence annotation coverage','Pfam annotations unavailable.');empty(bx,'Domain-family inventory','No families are inferred from parent names or structure appearance.')
    protocol_summaries=diversity.get('protocols',[])
    if protocol_summaries:
        candidate_counts=[];control_notes=[]
        valid_states={'qualified_domain_available','unclassified_failed_domain_confidence_or_length_gate','unclassified_no_Pfam_GA_span'}
        for i,entry in enumerate(protocol_summaries):
            pid=entry['protocol_id'];valid=[r for r in diversity_models if r.get('protocol_id')==pid and r.get('role')!='control' and r.get('analysis_status') in valid_states]
            actual=len(valid);classified=sum(num(r,'qualified_domain_count',0)>0 for r in valid);unclassified=actual-classified;candidate_counts.append(actual)
            cx.bar(i,unclassified,color=GRAY,label='Unclassified' if i==0 else None);cx.bar(i,classified,bottom=unclassified,color=BLUE,label='Qualified span' if i==0 else None)
            candidate_spans=[r for r in diversity_domains if r.get('protocol_id')==pid and r.get('role')!='control' and r.get('gate_status')=='qualified']
            candidate_clusters={r['cluster_id'] for r in candidate_spans if r.get('cluster_id')}
            cluster_label=f"\n{len(candidate_clusters)} candidate-associated cluster(s)" if candidate_clusters else ''
            cx.text(i,actual+.15,f"{classified}/{actual} candidates with qualified spans\n{unclassified} unclassified"+cluster_label,ha='center',fontsize=8)
            controls=[r for r in diversity_domains if r.get('protocol_id')==pid and r.get('role')=='control' and r.get('gate_status')=='qualified']
            if controls:control_notes.append(f"Controls: {len(controls)} spans on {len({r['peptide_id'] for r in controls})} sequence(s); {len({r['cluster_id'] for r in controls if r.get('cluster_id')})} internal cluster(s).")
        cx.set(xticks=range(len(protocol_summaries)),xticklabels=[protocol_label(p['protocol_id']).replace(' ','\n') for p in protocol_summaries],ylabel='Sequence-verified candidate models',title='Candidate 3D gate coverage');cx.set_ylim(0,max(candidate_counts+[0])*1.35+1)
        cx.legend(loc='upper left',fontsize=8)
        cx.text(0,-.25,'Gate: contiguous ≥50 aa; ≥80% residues pLDDT ≥70.\n'+('\n'.join(control_notes) if control_notes else 'No control cluster result reported.'),transform=cx.transAxes,size=8,color=GRAY)
    else:empty(cx,'3D gate coverage','No actual confidence-qualified domain audit is available. Structural clustering remains unavailable.')
    save(f)

    if region_peptides or paired_fragments or paired_native:
        f=figure('06_region_context_appendix','Appendix · identical residues in different sequence contexts',
            'analysis/region_peptide_summary.tsv; paired_controls.tsv; paired_native_regions.tsv',
            'Region fractions are averaged within peptide across preserved origin hypotheses. Comparisons retain deduplicated exact amino-acid intervals; paired dots can share peptides and are not independent biological replicates.',
            'Same-amino-acid comparisons are predictor-context differences, not biological effects. Short tails and mapping alternatives limit interpretation; marginal region distributions are not paired effects.')
        axs=f.subplots(1,3)
        for key,label,color in [('junction_window','Junction window',BLUE),('parent_a_canonical','Retained parent A',GRAY),('parent_b_out_of_frame','Out-of-frame parent B',RED)]:
            values=sorted(num(r,'mean_f_idr') for r in region_peptides if r.get('region')==key and num(r,'mean_f_idr') is not None)
            if values: axs[0].step(values,np.arange(1,len(values)+1)/len(values),where='post',color=color,label=f'{label} (n={len(values)})')
        axs[0].set(xlim=(0,1),ylim=(0,1.02),xlabel='Mean region f_IDR across ORF origins',ylabel='Fraction of eligible peptide regions',title='Region-specific sequence predictions');axs[0].legend(fontsize=8,loc='lower right')
        for ax,rows,key,title in [(axs[1],paired_fragments,'isolated_fragment_f_idr','Fusion vs isolated fragment'),(axs[2],paired_native,'native_parent_region_f_idr','Fusion vs native-parent context')]:
            if rows:
                ax.scatter([num(r,key) for r in rows],[num(r,'candidate_region_f_idr') for r in rows],s=20,alpha=.4,color=BLUE if key.startswith('isolated') else RED);ax.plot([0,1],[0,1],color=GRAY,ls=':');ax.set(xlim=(0,1),ylim=(0,1),xlabel='Same interval: comparison-context f_IDR',ylabel='Same interval: fusion-context f_IDR',title=title)
                ax.text(.03,.95,f'{len(rows)} interval comparisons\nMedian Δ = {np.median([num(r,"delta_f_idr") for r in rows]):.3f}',transform=ax.transAxes,va='top',fontsize=11)
            else: empty(ax,title,'Paired sequence context predictions unavailable.')
        save(f)

    parent_source=ROOT/'results/structures/gsdmd_parent'
    fusion_source=ROOT/'results/structures/gsdmd_tmem106a/cartoon.png'
    example_source=ROOT/'results/presentation/candidate_examples.json'
    if (parent_source/'cartoon.png').exists() and fusion_source.exists() and example_source.exists():
        assets=OUT/'assets';assets.mkdir(exist_ok=True)
        for source,dest in [(parent_source/'cartoon.png',assets/'experimental_parent_6N9N.png'),(parent_source/'6N9N.cif',assets/'experimental_parent_6N9N.cif'),(parent_source/'rcsb_entry.json',assets/'experimental_parent_6N9N_metadata.json'),(fusion_source,assets/'cached_gsdmd_tmem106a_prediction.png')]:
            INPUTS[str(source.relative_to(ROOT))]=hashlib.sha256(source.read_bytes()).hexdigest();shutil.copyfile(source,dest)
        INPUTS[str(example_source.relative_to(ROOT))]=hashlib.sha256(example_source.read_bytes()).hexdigest()
        example=next(c for c in json.loads(example_source.read_text())['candidates'] if c['pair_id']=='Gsdmd:Tmem106a')
        (assets/'gsdmd_rna_example.json').write_text(json.dumps({'source':str(example_source.relative_to(ROOT)),'source_sha256':INPUTS[str(example_source.relative_to(ROOT))],'example':example},indent=2))
        f=figure('07_experimental_parent_context','Appendix · experimental parent and conditional fusion context',
            'RCSB PDB 6N9N; exact probe/GENCODE M28 transcript mapping; sequence-verified cached Boltz2 Gsdmd:Tmem106a model',
            'The experimental mouse GSDMD parent is an X-ray structure. RNA exons are a representative exact-matching transcript pair, drawn with symbolic widths. The 118-aa fusion is a reference-reconstructed hypothesis with a cached MSA-backed prediction.',
            'An experimental parent structure does not validate the fusion fold. RNA exon provenance, translation, confidence and function remain separate evidence.')
        left,middle,right=f.subplots(1,3,gridspec_kw={'width_ratios':[1,1.2,1]})
        left.imshow(plt.imread(assets/'experimental_parent_6N9N.png'));left.axis('off');left.set_title('Experimental mouse GSDMD\nPDB 6N9N · X-ray · 3.30 Å',fontsize=13)
        right.imshow(plt.imread(assets/'cached_gsdmd_tmem106a_prediction.png'));right.axis('off');right.set_title('Predicted conditional fusion\nCached MSA-backed Boltz2\nMean pLDDT 48.7',fontsize=13)
        middle.axis('off');middle.set_title('Conditional RNA and protein hypothesis',fontsize=13)
        a=example['probe_examples'][0]['sides']['a']['representative'];b=example['probe_examples'][0]['sides']['b']['representative']
        middle.text(.02,.85,'Gsdmd',color=BLUE,weight='bold',transform=middle.transAxes);middle.text(.51,.85,'Tmem106a',color=RED,weight='bold',transform=middle.transAxes)
        for x,n in [(.02,1),(.20,2)]:
            middle.add_patch(plt.Rectangle((x,.69),.14,.09,color=BLUE,transform=middle.transAxes));middle.text(x+.07,.735,str(n),color='white',ha='center',va='center',transform=middle.transAxes)
        for x,n in [(.50,6),(.63,7),(.76,8),(.89,9)]:
            middle.add_patch(plt.Rectangle((x,.69),.1,.09,color=RED,transform=middle.transAxes));middle.text(x+.05,.735,str(n),color='white',ha='center',va='center',transform=middle.transAxes)
        middle.annotate('',xy=(.48,.735),xytext=(.36,.735),xycoords='axes fraction',arrowprops={'arrowstyle':'->','color':INK})
        middle.text(.02,.58,f"{a['chromosome']}:{a['breakpoint']:,} →\n{b['chromosome']}:{b['breakpoint']:,}\nGRCm39; exon widths symbolic",transform=middle.transAxes,fontsize=10)
        cut=73/118;middle.add_patch(plt.Rectangle((.02,.34),.96*cut,.09,color=BLUE,transform=middle.transAxes));middle.add_patch(plt.Rectangle((.02+.96*cut,.34),.96*(1-cut),.09,color=RED,transform=middle.transAxes))
        middle.text(.04,.37,'1–73',color='white',transform=middle.transAxes);middle.text(.02+.96*cut+.02,.37,'74–118',color='white',transform=middle.transAxes)
        middle.text(.02,.20,'118 aa · blue: parent-A RNA origin\nRed: parent-B RNA / split codon 74\nConditional transcript chain; not isoform validation',transform=middle.transAxes,fontsize=10)
        save(f)

    single_assets=[a for a in structure_assets if 'single_sequence' in a.get('protocol_id','') and a.get('png') and (ROOT/a['png']).exists()]
    single_assets=single_assets[:12]
    for offset in range(0,len(single_assets),6):
        block=single_assets[offset:offset+6]
        f=figure(f'{8+offset//6:02}_single_sequence_structure_gallery','Appendix · actual single-sequence candidate predictions',
            'compute/structure_gallery.json; sequence-verified model artifacts and audited source maps',
            'Display up to the first 12 available candidate models in frozen manifest order, six per sheet, without filtering on confidence. All available candidate cartoons remain in the offline atlas. Blue denotes RNA parent A; red parent B or split codon. Single-sequence Boltz2 protocol is separate from the cached MSA reference.',
            'These are conditional sequence hypotheses and predicted geometries, not experimental folds. Low confidence remains visible; appearance is not a function score.')
        axes=f.subplots(2,3)
        for ax,a in zip(axes.flat,block):
            path=ROOT/a['png'];INPUTS[str(path.relative_to(ROOT))]=hashlib.sha256(path.read_bytes()).hexdigest();ax.imshow(plt.imread(path));ax.axis('off');ax.set_title(a.get('title','Candidate')+f" · pLDDT {a.get('mean_plddt',float('nan')):.1f}",fontsize=10)
        for ax in list(axes.flat)[len(block):]:ax.axis('off')
        save(f)

    cross_model_comparison()

    for relative,digest in INPUTS.items():
        if hashlib.sha256((ROOT/relative).read_bytes()).hexdigest()!=digest: raise RuntimeError('Input changed during rendering; rerun after snapshot freezes: '+relative)
    headline='Conditional protein hypotheses are usually not predominantly disordered' if conditional and sum(num(r,'f_idr',0)>.5 for r in conditional)<len(conditional)/2 else 'Conditional protein disorder depends on sequence and predictor'
    manifest={'headline':headline,'renderer_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'status':'available_inputs_rendered','inputs':INPUTS,'figures':FIGURES,'n_disorder_rows':len(disorder),'n_model_rows':len(metrics),'n_all_model_records':len(all_metrics),'cohort_tiers':tiers,'limitations':['RNA support is not full-transcript or translation validation.','Conditional and alternative ORF hypotheses are separate from strict-primary eligibility.','Predicted disorder and low structural confidence are distinct.','Missing controls or independent predictors are not inferred.'],'build_command':'PYTHONPATH=/tmp/chrna-presentation-tools .venv/bin/python scripts/structure_campaign/render_report.py'}
    (OUT/'manifest.json').write_text(json.dumps(manifest,indent=2))
    (OUT/'input_snapshot.json').write_text(json.dumps({'analysis_summary':summary,'novel_tail_summary':novel_tail_summary,'census_summary':census_summary,'region_summary':region_summary,'alternative_start_summary':alternative_summary,'cohort_summary':audit,'compute_jobs':jobs,'compute_summary':compute_summary,'compute_receipt':compute_receipt,'cleanup_receipt':cleanup_receipt,'domain_manifest':domain_manifest,'diversity_summary':diversity},indent=2))
    sections=[]
    for item in FIGURES:
        paths={ext:html.escape(os.path.relpath(ROOT/item[ext],OUT)) for ext in ('svg','png','pdf')}
        sections.append(f'<section id="{item["slug"]}"><h2>{html.escape(item["title"])}</h2><img src="{paths["svg"]}" alt="{html.escape(item["title"])}"><p><b>Method.</b> {html.escape(item["method"])}</p><p><b>Interpretation.</b> {html.escape(item["interpretation"])}</p><small>Source: {html.escape(item["source"])}</small><p><a href="{paths["png"]}">PNG</a> · <a href="{paths["svg"]}">SVG</a> · <a href="{paths["pdf"]}">PDF</a></p></section>')
    for item in FIGURES:
        if item.get('related_report'):
            sections.append('<section><p><a href="'+html.escape(os.path.relpath(ROOT/item['related_report'],OUT))+'">Open the exact-sequence cross-model comparison and actual coordinate/confidence artifacts</a></p></section>')
    page='''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>chRNA disorder / structure atlas</title><style>body{font:17px/1.55 system-ui,sans-serif;color:#172431;max-width:1500px;margin:0 auto;padding:32px;background:#f5f6fa}header,section{background:white;padding:28px;margin:20px 0;border:1px solid #dce0e7;border-radius:12px}header{border-top:6px solid #0000ff}h1{font-size:38px}h2{font-size:28px}img{width:100%;height:auto}a{color:#0000ff}small{color:#63707d}.note{border-left:5px solid #d22d27;padding:12px;background:#fceeed}@media(max-width:600px){body{padding:12px}header,section{padding:16px}h1{font-size:30px;line-height:1.2;overflow-wrap:anywhere}h2{font-size:24px}small{overflow-wrap:anywhere}}</style></head><body><header><h1>chRNA disorder / structure atlas</h1><p>Source-bound sequence predictions and structural coverage. This report reads actual cached outputs and keeps unavailable analyses visible.</p><p class="note">Reported RNA support does not establish a full transcript, translated protein, folded structure or biological function. Strict-primary and conditional ORF hypotheses remain separate.</p><p><a href="structure_campaign_gallery.pptx">PowerPoint gallery</a> · <a href="manifest.json">Figure manifest and input hashes</a> · <a href="input_snapshot.json">Analysis status snapshot</a></p></header>'''+''.join(sections)+'</body></html>'
    by_model={}
    for r in metrics:
        if r.get('sequence_sha256'): by_model.setdefault(r['sequence_sha256'],[]).append(r)
    table_rows=[]
    for row in hypotheses:
        models=by_model.get(row.get('sequence_sha256'),[])
        model_values=[{'protocol_id':m.get('protocol_id'),'mean_plddt':num(m,'mean_plddt'),'fraction_plddt_ge70':num(m,'fraction_plddt_ge70')} for m in models]
        table_rows.append({'peptide_id':row.get('peptide_id'),'pair_ids':row.get('pair_ids'),'tier':row.get('cohort_tier'),'length':num(row,'length'), 'f_idr':num(row,'f_idr'),'v1_f_idr':num(row,'v1_f_idr'),'models_by_protocol':model_values})
    (OUT/'peptide_atlas.json').write_text(json.dumps(table_rows,indent=2))
    rows_html=[]
    for row in table_rows:
        values=[row['pair_ids'],clean_tier(row['tier']),str(int(row['length'])) if row['length'] is not None else 'Unknown',f"{row['f_idr']:.3f}" if row['f_idr'] is not None else 'Unavailable',f"{row['v1_f_idr']:.3f}" if row['v1_f_idr'] is not None else 'Unavailable','; '.join(m['protocol_id']+': '+str(round(m['mean_plddt'],1)) for m in row['models_by_protocol']) or 'Unavailable',row['peptide_id']]
        rows_html.append('<tr>'+''.join('<td>'+html.escape(str(v or 'Unknown'))+'</td>' for v in values)+'</tr>')
    table='<section><h2>Searchable peptide atlas</h2><p>Each row is a sequence hypothesis. A missing model remains unavailable. Exact peptide IDs preserve the connection to the source tables.</p><input id="search" placeholder="Filter by ordered pair, tier or peptide ID" style="padding:12px;width:90%;font:inherit"><p id="count"></p><div style="overflow:auto;max-height:650px"><table id="atlas"><thead><tr>'+''.join('<th>'+h+'</th>' for h in ['Ordered pairs','Eligibility tier','Length','V3 f_IDR','V1 f_IDR','pLDDT by protocol','Peptide ID'])+'</tr></thead><tbody>'+''.join(rows_html)+'</tbody></table></div><p><a href="peptide_atlas.json">Download displayed atlas rows</a> · <a href="../analysis/candidate_evidence_table.tsv">Full joined evidence TSV and protein-evidence status</a></p></section>'
    table+='<style>table{border-collapse:collapse;font-size:14px}th,td{padding:10px;border-bottom:1px solid #ddd;text-align:left}th{position:sticky;top:0;background:#edf0f8}td:last-child{font:10px monospace;max-width:190px;overflow-wrap:anywhere}</style><script>const rows=[...document.querySelectorAll("#atlas tbody tr")];function filter(){const q=document.querySelector("#search").value.toLowerCase();let n=0;rows.forEach(r=>{r.hidden=!r.textContent.toLowerCase().includes(q);if(!r.hidden)n++});document.querySelector("#count").textContent=n+" / "+rows.length+" sequence records"}document.querySelector("#search").addEventListener("input",filter);filter();</script>'
    if audit.get('counts'):
        c=audit['counts'];scope=f"Audited scope: {c.get('primary_reference_pairs')} RNA pairs; {c.get('strict_primary_peptides')} strict-primary peptides; {c.get('conditional_census_unique_peptides')} conditional annotated-start peptides from {c.get('conditional_census_pairs')} pairs. The {c.get('alternative_start_only_unique_peptides')} alternative-start hypotheses remain a separate sensitivity tier."
        page=page.replace('</header>','<p class="note">'+html.escape(scope)+'</p></header>')
    if conditional:
        values=[num(r,'f_idr') for r in conditional];n3=sum(v>.5 for v in values);n1=sum(num(r,'v1_f_idr',0)>.5 for r in conditional)
        lengths=[num(r,'length',0) for r in conditional];weighted=sum(v*l for v,l in zip(values,lengths))/sum(lengths)
        quantitative=f"Among {len(conditional)} conditional annotated-start peptide hypotheses, V3 predicts {n3} ({n3/len(conditional):.1%}) predominantly disordered; related V1 predicts {n1} ({n1/len(conditional):.1%}). V3 mean f_IDR is {np.mean(values):.3f}, median {np.median(values):.3f}, and residue-weighted mean {weighted:.3f}. These computed census fractions do not estimate nonfunction or experimentally measured disorder."
        page=page.replace('</header>','<p>'+html.escape(quantitative)+'</p></header>')
        (OUT/'quantitative_summary.json').write_text(json.dumps({'scope':'conditional_annotated_start_hypotheses_not_observed_proteins','n':len(conditional),'v3_n_predominantly_disordered':n3,'v1_n_predominantly_disordered':n1,'v3_mean_f_idr':float(np.mean(values)),'v3_median_f_idr':float(np.median(values)),'v3_residue_weighted_f_idr':weighted,'consensus_subset_n':len(consensus)},indent=2))
    alt=alternative_summary.get('groups',{}).get('candidate:alternative_start_sensitivity')
    if alt:
        alt_text=f"Separate alternative-start sensitivity: {alt['n_predominantly_disordered']}/{alt['n_peptides']} ({alt['fraction_predominantly_disordered']:.1%}) predominantly disordered by V3; V1 {alt['v1_fraction_predominantly_disordered']:.1%}. These hypotheses are outside the 188-peptide annotated-start census."
        page=page.replace('</header>','<p>'+html.escape(alt_text)+'</p></header>')
    cards=[]
    for asset in structure_assets:
        if asset.get('png') and (ROOT/asset['png']).exists():
            local=os.path.relpath(ROOT/asset['png'],OUT)
            cards.append('<article><h3>'+html.escape(asset.get('title','Candidate'))+'</h3><p>'+html.escape(asset.get('protocol_id','Protocol unspecified'))+' · '+('Cached prior model' if asset.get('reused_prior_prediction') else 'New prediction')+'</p><a href="'+html.escape(local)+'"><img src="'+html.escape(local)+'" alt="'+html.escape(asset.get('title','Predicted structure'))+'"></a><p>'+html.escape(asset.get('confidence_status','Confidence status unavailable'))+'</p></article>')
    gallery='<section><h2>Actual available candidate structure gallery</h2><p>Audited RNA-source colors; protocols and confidence remain explicit. Model appearance is not experimental validation.</p><div style="display:grid;grid-template-columns:repeat(auto-fit,minmax(300px,1fr));gap:24px">'+''.join(cards)+'</div></section>' if cards else ''
    page=page.replace('<h1>chRNA disorder / structure atlas</h1>','<h1>'+html.escape(headline)+'</h1>')
    page=page.replace('</header>','<p><a href="../RESULTS.md">Full results and interpretation</a> · <a href="../EXECUTION.md">Execution record and limits</a></p></header>')
    if compute_receipt:
        cost=compute_receipt.get('cleanup_and_cost',{})
        execution=f"Actual GPU execution: {compute_receipt.get('new_models_total','?')} new Boltz2 models across {compute_receipt.get('unique_new_modeled_sequences','?')} sequences, including {compute_receipt.get('new_diagnostic_seed_models','?')} extra seed models. Cached MSA context is separate. GPU status: {cost.get('status','unknown').replace('_',' ')}."
        if num(cost,'gpu_elapsed_cost_estimate_usd') is not None:
            execution+=f" Quoted-rate GPU elapsed estimate ${num(cost,'gpu_elapsed_cost_estimate_usd'):.2f}; not a provider invoice or total storage/egress cost."
        page=page.replace('</header>','<p>'+html.escape(execution)+' <a href="../compute/campaign_receipt.json">Execution and cost receipt</a></p></header>')

    ledger=[];ledger_html=[]
    for job in jobs if isinstance(jobs,list) else jobs.get('jobs',[]):
        validation=job.get('validation',{});available_files={}
        for name,relative in job.get('artifacts',{}).items():
            target=(ROOT/relative).resolve()
            if target.is_relative_to(BASE.resolve()) and target.exists():available_files[name]=os.path.relpath(target,OUT)
        record={'job_id':job.get('job_id'),'peptide_id':job.get('peptide_id'),'sequence_sha256':job.get('sequence_sha256'),'pair_ids':job.get('input_mapping',{}).get('pair_ids',''),'role':job.get('role'),'protocol_id':job.get('protocol_id'),'seed':job.get('seed'),'status':job.get('status'),'length_aa':job.get('length_aa'),'mean_plddt':validation.get('mean_plddt'),'files':available_files,'artifact_sha256':job.get('artifact_sha256',{})}
        ledger.append(record)
        values=[record['pair_ids'] or str(record['peptide_id'])[:24],record['role'],record['protocol_id'],record['seed'],record['status'],record['length_aa'],f"{record['mean_plddt']:.1f}" if record['mean_plddt'] is not None else 'Unavailable']
        links=' · '.join('<a href="'+html.escape(available_files[key])+'">'+label+'</a>' for key,label in [('model.cif','Coordinates'),('residue_confidence.tsv','Residue confidence'),('pae.npz','PAE')] if key in available_files)
        ledger_html.append('<tr>'+''.join('<td>'+html.escape(str(v))+'</td>' for v in values)+'<td>'+ (links or 'Unavailable')+'</td></tr>')
    (OUT/'model_ledger.json').write_text(json.dumps(ledger,indent=2))
    model_table='<section id="models"><h2>Every planned and executed model</h2><p>One row per execution job, including controls and deferred runs. Repeated seeds are repeated models, not additional proteins. Protocols stay separate. Download actual coordinates, residue confidence and PAE when available.</p><input id="model-search" placeholder="Filter by pair, role, protocol, seed or status" style="padding:12px;width:90%;font:inherit"><p id="model-count"></p><div style="overflow:auto;max-height:600px"><table id="model-table"><thead><tr>'+''.join('<th>'+h+'</th>' for h in ['Pair / control','Role','Protocol','Seed','Status','Length','Mean pLDDT','Artifacts'])+'</tr></thead><tbody>'+''.join(ledger_html)+'</tbody></table></div><p><a href="model_ledger.json">Download displayed model ledger</a> · <a href="../compute/jobs.json">Full job provenance</a></p></section><script>const modelRows=[...document.querySelectorAll("#model-table tbody tr")];function filterModels(){const terms=document.querySelector("#model-search").value.toLowerCase().trim().split(/ +/);let n=0;modelRows.forEach(r=>{r.hidden=!terms.every(t=>r.textContent.toLowerCase().includes(t));if(!r.hidden)n++});document.querySelector("#model-count").textContent=n+" / "+modelRows.length+" execution records"}document.querySelector("#model-search").addEventListener("input",filterModels);filterModels();</script>'
    rarefaction_notes=[]
    for entry in diversity.get('protocols',[]):
        protocol=entry['protocol_id'];observed=[r for r in rarefaction if r.get('protocol_id')==protocol]
        if observed:
            endpoint=max(observed,key=lambda r:num(r,'sampled_actual_peptides',0))
            actual_candidates=[r for r in diversity_models if r.get('protocol_id')==protocol and r.get('role')!='control' and r.get('analysis_status') in {'qualified_domain_available','unclassified_failed_domain_confidence_or_length_gate','unclassified_no_Pfam_GA_span'}]
            covered=sum(num(r,'qualified_domain_count',0)>0 for r in actual_candidates)
            message=f"{protocol}: descriptive rarefaction was computed through {int(num(endpoint,'sampled_actual_peptides',0))} actual candidate peptides; endpoint {num(endpoint,'expected_observed_clusters',0):.2f} observed internal clusters. Only {covered}/{len(actual_candidates)} candidate models have qualified spans; {len(actual_candidates)-covered} remain unclassified."
            total=num(endpoint,'total_actual_candidate_peptides',0);richness=num(endpoint,'expected_observed_clusters',0)
            if total and all(abs(num(r,'expected_observed_clusters',0)-richness*num(r,'sampled_actual_peptides',0)/total)<1e-8 for r in observed):
                message+=f" The observed curve is the simple linear relation {richness:g} × sampled peptides / {total:g}; it does not establish population richness."
        else:
            message=f"{protocol}: no rarefaction curve displayed — {entry.get('rarefaction_status','unavailable').replace('_',' ')}."
        rarefaction_notes.append('<li>'+html.escape(message)+'</li>')
    rarefaction_section='<section><h2>Observed structural richness and unavailable analyses</h2><ul>'+''.join(rarefaction_notes)+'</ul><p>Only actual confidence-qualified spans enter internal clusters. The selected finite sample is not an estimate of all chRNA structural diversity. No external fold-novelty search or DSSP secondary-structure census was run.</p><p><a href="../diversity/rarefaction.tsv">Actual rarefaction table</a> · <a href="../diversity/manifest.json">Diversity methods, gates, and snapshot</a></p></section>'
    page=page.replace('</body>',rarefaction_section+gallery+model_table+table+'</body>')
    (OUT/'index.html').write_text(page)
    md=['# Campaign figure guide','', 'These figures are generated only from available audited inputs. Empty states are intentional.','']
    for item in FIGURES: md += [f"## {item['title']}",'',f"![{item['title']}]({os.path.relpath(ROOT/item['png'],OUT)})",'',f"Source: {item['source']}",'',f"Method: {item['method']}",'',f"Interpretation: {item['interpretation']}",'']
    (OUT/'FIGURE_GUIDE.md').write_text('\n'.join(md))
    write_notes(metrics,all_metrics,diversity_models,diversity_domains)
    return manifest

def write_notes(metrics,all_metrics,diversity_models,diversity_domains):
    single_first=[r for r in metrics if r.get('protocol_id','').endswith('single_sequence')]
    single_executed=[r for r in all_metrics if r.get('protocol_id','').endswith('single_sequence') and num(r,'mean_plddt') is not None]
    candidate_first=[r for r in single_first if r.get('role')!='control'];control_first=[r for r in single_first if r.get('role')=='control']
    n_extra=len(single_executed)-len(single_first)
    audited_candidates=[r for r in diversity_models if r.get('protocol_id','').endswith('single_sequence') and r.get('role')!='control' and r.get('analysis_status') in {'qualified_domain_available','unclassified_failed_domain_confidence_or_length_gate','unclassified_no_Pfam_GA_span'}]
    n_qualified=sum(num(r,'qualified_domain_count',0)>0 for r in audited_candidates)
    candidate_clusters={r['cluster_id'] for r in diversity_domains if r.get('protocol_id','').endswith('single_sequence') and r.get('role')!='control' and r.get('cluster_id')}
    notes=f"""# Five-minute campaign presentation

Counts below come from this build's saved source snapshot. Use the frozen final deck for presentation; execution and diversity snapshots may advance during collection.

- **0:00–0:25 — Slide 1: main result and workflow.** “Conditional protein hypotheses are usually not predominantly disordered by these predictors.” The primary V3 result is 42/188 (22.3%); related V1 sensitivity gives 18/188 (9.6%). These are sequence predictions, not stability or function measurements.
- **0:25–1:10 — Slide 2 / figure 01: define the denominator.** Start with 109 reported-supported RNA gene pairs. Reference reconstruction gives 188 annotated-start peptide hypotheses across 91 pairs; the 30 consensus peptides are a nested subset. Strict full-chain primary eligibility is zero. The 746 alternative-start hypotheses are separate sensitivity cases. Neither 109 nor 934 is the count of demonstrated proteins.
- **1:10–1:50 — Slide 3 / figure 02: show the full distribution.** Explain f_IDR as the predicted disordered residue fraction; “predominantly” means more than half of a peptide. The 22.3% peptide classification fraction differs from the 30.7% peptide-equal mean f_IDR. V1 is a related-network sensitivity analysis, not independent validation. No predictor-error confidence interval or universal chRNA prevalence is claimed.
- **1:50–2:30 — Slides 5 and 8 / figures 04 and 07: Gsdmd caution case.** V3 predicts f_IDR=1.0 and V1 about0.483 for the118-aa reference; the cached MSA-backed Boltz model has mean pLDDT48.7. The paper's functional exemplar shows why predicted disorder or low confidence cannot rule out function. The 6N9N X-ray structure is an experimental parent, not fusion ground truth. The exon diagram is a conditional reference reconstruction.
- **2:30–3:35 — Slides 4 and 6 / figures 03 and 05: actual NVIDIA execution and coverage.** Open-source Boltz2 2.2.1 ran on an NVIDIA A10080GB in an explicitly separate single-sequence arm. This snapshot has {len(single_first)} verified first-pass sequences ({len(candidate_first)} candidate hypotheses, {len(control_first)} controls) and {n_extra} additional seed models: {len(single_executed)} executed single-sequence models total. Repeats are not extra proteins. The prior cached MSA-backed reference remains separate; hosted NIM is not claimed. Pfam annotates151/188 sequences, but sequence families are not3D clusters. In the diversity snapshot, {n_qualified}/{len(audited_candidates)} candidate models pass the fixed domain gate, representing {len(candidate_clusters)} candidate-associated internal cluster(s). Controls, overlapping spans and unclassified models stay explicit; no novel-fold claim.
- **3:35–4:20 — Slide 7 / figure 06, then slides9–10 / figures08–09 where available.** Exact retained-region comparisons have median predicted differences of zero in both isolated-fragment and native-parent context. These are dependent computational intervals, not biological replicates. Show a few source-colored cartoons in fixed order; attractive geometry is not a function score. The offline atlas exposes every model job and its actual artifacts.
- **4:20–5:00 — Close on the evidence gap.** The next experiments must establish a particular full transcript, translation and a specified biological activity. These analyses prioritize conditional hypotheses and reveal model sensitivity. They do not classify chRNAs as functional/nonfunctional or establish broad prevalence beyond the reconstructed reference.

Detailed sources and limits: [Results](../RESULTS.md), [Execution](../EXECUTION.md), [figure guide](FIGURE_GUIDE.md), [input hashes](manifest.json).
"""
    notes=notes.replace('about0.483','about 0.483').replace('the118-aa','the 118-aa').replace('pLDDT48.7','pLDDT 48.7').replace('A10080GB','A100 80 GB').replace('annotates151','annotates 151').replace('not3D','not 3D').replace('slides9–10','slides 9–10').replace('figures08–09','figures 08–09')
    cross_path=BASE/'cross_model_gsdmd/report/comparison_snapshot.json'
    if cross_path.exists() and any(f['slug']=='10_cross_model_gsdmd' for f in FIGURES):
        cross=json.loads(cross_path.read_text());actual=[m for m in cross.get('models',[]) if m.get('status') in {'verified','cached_verified','completed'}];new=[m for m in actual if not m.get('reused_prior_prediction')]
        means='; '.join(m.get('engine','Model')+f" mean pLDDT {num(m,'mean_plddt'):.1f}" for m in new if num(m,'mean_plddt') is not None)
        methods=cross.get('disorder_methods',[]);regions='; '.join(m['method']+f" donor {num(m,'donor_fraction_disordered'):.1%}, tail {num(m,'tail_fraction_disordered'):.1%}" for m in methods if num(m,'donor_fraction_disordered') is not None and num(m,'tail_fraction_disordered') is not None)
        notes+='\n## Optional 45-second cross-model comparison — slide 11\n\n'
        notes+=f"The same exact 118-aa sequence has {len(actual)} audited model records, including {len(new)} new predictions in this separate comparison. {means}. The Boltz MSA-backed baseline and three sequence-only seeds are reused prior predictions, not new refolds.\n\n"
        notes+=regions+'. These are sequence-disorder predictions, not disorder fractions inferred from folding confidence. V1/V3 are related networks. The main disagreement concerns the retained donor segment; cartoons use independent orientations, not a structural superposition. Computational consistency does not establish the physical fusion fold or function. The original campaign census and 67-model execution denominator remain unchanged.\n'
    (OUT/'PRESENTATION_NOTES.md').write_text(notes)

def presentation(manifest):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    def text(slide,x,y,w,h,value,size=20,color=INK,bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
        tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
        p=tf.paragraphs[0];p.text=value;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color[1:]);return box
    slide=prs.slides.add_slide(prs.slide_layouts[6]);text(slide,.6,.5,14.8,1,manifest['headline'],32,BLUE,True)
    q=json.loads((OUT/'quantitative_summary.json').read_text()) if (OUT/'quantitative_summary.json').exists() else {}
    if q: text(slide,.6,1.65,14.8,.6,f"V3 {q['v3_n_predominantly_disordered']/q['n']:.1%} · V1 {q['v1_n_predominantly_disordered']/q['n']:.1%} · {q['n']} conditional peptide hypotheses",25,GRAY)
    steps=['Audit RNA evidence','Resolve complete ORFs','Predict sequence disorder','Model selected structures','Measure coverage / diversity']
    for i,s in enumerate(steps):
        x=.6+i*3.05;text(slide,x,2.6,2.7,.7,f'{i+1:02}',42,BLUE if i%2==0 else RED,True);text(slide,x,3.35,2.7,1.6,s,25,INK,True)
    text(slide,.6,5.6,14.5,1.2,'Strict-primary, conditional-reference and alternative ORF tiers stay separate. Missing measurements remain missing.',29,RED,True)
    text(slide,.6,7.2,14.5,.9,'Offline atlas · five source-bound figures · sequence disorder ≠ low fold confidence ≠ nonfunction',22,GRAY)
    slide.notes_slide.notes_text_frame.text='Conceptual workflow. Source: plans/chRNA_disorder_diversity_campaign.md. It is not a claim that every stage completed. Exact execution status and input hashes are in report/manifest.json.'
    for item in FIGURES:
        slide=prs.slides.add_slide(prs.slide_layouts[6]);text(slide,.5,.2,15,.7,item['title'],27,INK,True)
        slide.shapes.add_picture(str(ROOT/item['slide_png']),Inches(.4),Inches(1),width=Inches(15.2),height=Inches(7.1))
        text(slide,.6,8.18,14.8,.65,item['interpretation'],13,GRAY)
        slide.notes_slide.notes_text_frame.text='SOURCE\n'+item['source']+'\n\nMETHOD\n'+item['method']+'\n\nINTERPRETATION\n'+item['interpretation']
    prs.save(OUT/'structure_campaign_gallery.pptx')

def render_pptx():
    lo=os.environ.get('CHRNA_LIBREOFFICE_BIN') or shutil.which('libreoffice');ppm=os.environ.get('CHRNA_PDFTOPPM_BIN') or shutil.which('pdftoppm')
    if not lo or not ppm: raise SystemExit('Set CHRNA_LIBREOFFICE_BIN and CHRNA_PDFTOPPM_BIN to render the actual PowerPoint.')
    env=os.environ.copy();env['SAL_USE_VCLPLUGIN']='svp'
    if os.environ.get('CHRNA_RENDER_LIBRARY_PATH'): env['LD_LIBRARY_PATH']=os.environ['CHRNA_RENDER_LIBRARY_PATH']
    profile=Path(tempfile.mkdtemp(prefix='chrna-campaign-lo-'))
    for _ in range(3):
        r=subprocess.run([lo,f'-env:UserInstallation={profile.as_uri()}','--headless','--norestore','--nofirststartwizard','--convert-to','pdf','--outdir',str(OUT),str(OUT/'structure_campaign_gallery.pptx')],env=env,capture_output=True,text=True,timeout=120)
        if r.returncode==0:break
        if r.returncode!=81:raise RuntimeError(r.stdout+r.stderr)
    dest=OUT/'pptx_render';dest.mkdir(exist_ok=True)
    for old in dest.glob('slide-*.png'): old.unlink()
    subprocess.run([ppm,'-png','-r','90',str(OUT/'structure_campaign_gallery.pdf'),str(dest/'slide')],env=env,check=True,timeout=120)
    from PIL import Image,ImageDraw
    from pypdf import PdfReader
    from pptx import Presentation
    expected=len(Presentation(OUT/'structure_campaign_gallery.pptx').slides)
    assert len(PdfReader(OUT/'structure_campaign_gallery.pdf').pages)==expected
    files=sorted(dest.glob('slide-*.png'));assert len(files)==expected
    canvas=Image.new('RGB',(1500,330*math.ceil(len(files)/3)),'#e7eaf0');draw=ImageDraw.Draw(canvas)
    for i,path in enumerate(files):
        im=Image.open(path).convert('RGB');im.thumbnail((480,270));x=10+(i%3)*500;y=10+(i//3)*330;canvas.paste(im,(x,y));draw.text((x,y+280),f'Slide {i+1} · actual PPTX render',fill='black')
    canvas.save(OUT/'gallery_contact_sheet.png')
    (OUT/'render_validation.json').write_text(json.dumps({'status':'passed','slides':len(files),'pdf_pages':expected,'method':'Actual PowerPoint → LibreOffice PDF → pdftoppm PNG','pptx_sha256':hashlib.sha256((OUT/'structure_campaign_gallery.pptx').read_bytes()).hexdigest(),'pdf_sha256':hashlib.sha256((OUT/'structure_campaign_gallery.pdf').read_bytes()).hexdigest()},indent=2))

def check_browser():
    from playwright.sync_api import sync_playwright
    manifest=json.loads((OUT/'manifest.json').read_text());rows=json.loads((OUT/'peptide_atlas.json').read_text());errors=[]
    with sync_playwright() as runtime:
        browser=runtime.chromium.launch(headless=True,args=['--no-sandbox'])
        page=browser.new_page(viewport={'width':1440,'height':1000});page.on('pageerror',lambda e:errors.append(str(e)));page.goto((OUT/'index.html').as_uri())
        page.wait_for_function('Array.from(document.images).every(i=>i.complete && i.naturalWidth>0)')
        assert page.locator('#count').inner_text()==f'{len(rows)} / {len(rows)} sequence records'
        assert page.locator('img').count()>=len(manifest['figures'])
        ledger=json.loads((OUT/'model_ledger.json').read_text());assert page.locator('#model-count').inner_text()==f'{len(ledger)} / {len(ledger)} execution records'
        page.locator('#model-search').fill('single_sequence control');assert page.locator('#model-table tbody tr:visible').count()>=1;page.locator('#model-search').fill('')
        if rows:
            page.locator('#search').fill(rows[0]['pair_ids'].split(';')[0]);assert page.locator('#atlas tbody tr:visible').count()>=1;page.locator('#search').fill('')
        page.screenshot(path=str(OUT/'atlas_desktop.png'),full_page=True);page.set_viewport_size({'width':390,'height':844});assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth'), 'Mobile layout exceeds viewport';page.screenshot(path=str(OUT/'atlas_mobile.png'),full_page=True)
        assert not errors
        result={'status':'passed','offline_file_url':True,'figure_images':len(manifest['figures']),'total_images':page.locator('img').count(),'table_rows':len(rows),'model_execution_rows':len(ledger),'checks':['all local figures and actual structure assets loaded','real peptide-row count','gene-pair search filters results','desktop/mobile screenshots and mobile horizontal containment','no JavaScript errors'],'javascript_errors':errors,'html_sha256':hashlib.sha256((OUT/'index.html').read_bytes()).hexdigest()}
        cross=BASE/'cross_model_gsdmd/report';cross_page_path=cross/'index.html'
        if cross_page_path.exists() and any(f['slug']=='10_cross_model_gsdmd' for f in manifest['figures']):
            cross_page=browser.new_page(viewport={'width':1440,'height':1000});cross_page.on('pageerror',lambda e:errors.append(str(e)));cross_page.goto(cross_page_path.as_uri());cross_page.wait_for_function('Array.from(document.images).every(i=>i.complete && i.naturalWidth>0)')
            from urllib.parse import urlparse,unquote
            hrefs=cross_page.locator('a').evaluate_all('(links)=>links.map(a=>a.href)')
            for href in hrefs:
                parsed=urlparse(href)
                if parsed.scheme=='file': assert Path(unquote(parsed.path)).is_file(),href
            cross_page.screenshot(path=str(cross/'comparison_desktop.png'),full_page=True);cross_page.set_viewport_size({'width':390,'height':844});assert cross_page.evaluate('document.documentElement.scrollWidth <= window.innerWidth');cross_page.screenshot(path=str(cross/'comparison_mobile.png'),full_page=True)
            snapshot=json.loads((cross/'comparison_snapshot.json').read_text());actual=[m for m in snapshot['models'] if m.get('status') in {'verified','cached_verified','completed'}]
            assert all(m.get('sequence_verified') for m in actual);assert not errors
            (cross/'browser_validation.json').write_text(json.dumps({'status':'passed','offline_file_url':True,'verified_models':len(actual),'disorder_methods':len(snapshot.get('disorder_methods',snapshot.get('disorder_predictors',[]))),'local_links_checked':len(hrefs),'mobile_horizontal_containment':True,'javascript_errors':errors,'html_sha256':hashlib.sha256(cross_page_path.read_bytes()).hexdigest()},indent=2))
            result['cross_model_page_checked']=True
        browser.close();(OUT/'browser_validation.json').write_text(json.dumps(result,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--render-pptx',action='store_true');parser.add_argument('--check-browser',action='store_true');parser.add_argument('--cross-model-only',action='store_true');args=parser.parse_args()
    if args.cross_model_only:
        cross_model_comparison();print('Rendered separate exact-sequence cross-model comparison');raise SystemExit(0)
    manifest=render();presentation(manifest)
    if args.render_pptx: render_pptx()
    if args.check_browser: check_browser()
    print(json.dumps({'figures':len(FIGURES),'disorder_rows':manifest['n_disorder_rows'],'model_rows':manifest['n_model_rows'],'report':str(OUT/'index.html')}))
